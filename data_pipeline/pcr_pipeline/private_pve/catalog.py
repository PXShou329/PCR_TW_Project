from __future__ import annotations

import copy
import hashlib
import json
import re
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable, Mapping

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .media import MediaCatalog
from .models import SCHEMA_VERSION, WorkbookLayoutError


LOCAL_CATALOG_SCHEMA_VERSION = "private-pve-local-catalog/v1"
CATALOG_BUILDER_NAME = "pcr-private-pve-catalog"
CATALOG_BUILDER_VERSION = "1.0.0"

_SHA256 = re.compile(r"[0-9a-f]{64}\Z")
_MIME_EXTENSIONS = {
    "image/gif": "gif",
    "image/jpeg": "jpg",
    "image/png": "png",
}
_TOP_LEVEL_KEYS = {
    "assets",
    "diagnostics",
    "parser",
    "schema_version",
    "sheets",
    "source_workbook",
    "summary",
}
_ASSET_KEYS = {
    "byte_length",
    "file_extension",
    "height_px",
    "mime_type",
    "occurrences",
    "sha256",
    "width_px",
}
_OCCURRENCE_KEYS = {
    "anchor_cell",
    "column",
    "height_px",
    "row",
    "sheet_name",
    "width_px",
}
_ASSET_METADATA_KEYS = (
    "byte_length",
    "mime_type",
    "file_extension",
    "width_px",
    "height_px",
)


class CatalogMergeError(WorkbookLayoutError):
    """Raised when a staging catalog cannot be safely validated or merged."""


@dataclass(frozen=True)
class LocalCatalog:
    schema_version: str
    builder: dict[str, str]
    sources: tuple[dict[str, Any], ...]
    assets: tuple[dict[str, Any], ...]
    summary: dict[str, int]

    def to_dict(self) -> dict[str, Any]:
        return {
            "assets": copy.deepcopy(list(self.assets)),
            "builder": dict(self.builder),
            "schema_version": self.schema_version,
            "sources": copy.deepcopy(list(self.sources)),
            "summary": dict(self.summary),
        }

    def canonical_json_bytes(self) -> bytes:
        return _canonical_json_bytes(self.to_dict())


@dataclass(frozen=True)
class _ValidatedStaging:
    document: dict[str, Any]
    source_sha256: str
    semantic_sha256: str
    assets_by_sha: dict[str, dict[str, Any]]


def _canonical_json_bytes(value: Any) -> bytes:
    return (
        json.dumps(
            value,
            allow_nan=False,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        )
        + "\n"
    ).encode("utf-8")


def _object_without_duplicate_keys(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for key, value in pairs:
        if key in result:
            raise CatalogMergeError(f"duplicate JSON object key: {key!r}")
        result[key] = value
    return result


def _reject_json_constant(value: str) -> None:
    raise CatalogMergeError(f"non-finite JSON number is not allowed: {value}")


def _is_positive_int(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) and value > 0


def _require_sha256(value: Any, *, field: str) -> str:
    if not isinstance(value, str) or _SHA256.fullmatch(value) is None:
        raise CatalogMergeError(f"{field} must be a lowercase SHA-256 digest")
    return value


def _validate_source_workbook(value: Any, *, context: str) -> dict[str, Any]:
    if not isinstance(value, dict) or set(value) != {
        "byte_length",
        "filename",
        "sha256",
    }:
        raise CatalogMergeError(f"{context}.source_workbook has an unsupported shape")
    filename = value["filename"]
    if (
        not isinstance(filename, str)
        or not filename
        or Path(filename).name != filename
        or "/" in filename
        or "\\" in filename
    ):
        raise CatalogMergeError(
            f"{context}.source_workbook.filename must be a basename"
        )
    _require_sha256(value["sha256"], field=f"{context}.source_workbook.sha256")
    if not _is_positive_int(value["byte_length"]):
        raise CatalogMergeError(
            f"{context}.source_workbook.byte_length must be positive"
        )
    return value


def _validate_occurrence(
    value: Any,
    *,
    context: str,
    sheet_names: set[str],
) -> tuple[str, int, int, str, int, int]:
    if not isinstance(value, dict) or set(value) != _OCCURRENCE_KEYS:
        raise CatalogMergeError(f"{context} has an unsupported shape")
    sheet_name = value["sheet_name"]
    if not isinstance(sheet_name, str) or sheet_name not in sheet_names:
        raise CatalogMergeError(f"{context}.sheet_name is not a staged sheet")
    for field in ("row", "column", "width_px", "height_px"):
        if not _is_positive_int(value[field]):
            raise CatalogMergeError(f"{context}.{field} must be positive")
    expected_anchor = f"{get_column_letter(value['column'])}{value['row']}"
    if value["anchor_cell"] != expected_anchor:
        raise CatalogMergeError(
            f"{context}.anchor_cell does not match its row and column"
        )
    return (
        sheet_name,
        value["row"],
        value["column"],
        value["anchor_cell"],
        value["width_px"],
        value["height_px"],
    )


def _validate_asset(
    value: Any,
    *,
    context: str,
    sheet_names: set[str],
) -> dict[str, Any]:
    if not isinstance(value, dict) or set(value) != _ASSET_KEYS:
        raise CatalogMergeError(f"{context} has an unsupported shape")
    digest = _require_sha256(value["sha256"], field=f"{context}.sha256")
    for field in ("byte_length", "width_px", "height_px"):
        if not _is_positive_int(value[field]):
            raise CatalogMergeError(f"{context}.{field} must be positive")
    mime_type = value["mime_type"]
    extension = value["file_extension"]
    if _MIME_EXTENSIONS.get(mime_type) != extension:
        raise CatalogMergeError(
            f"{context} has conflicting MIME type and file extension"
        )
    occurrences = value["occurrences"]
    if not isinstance(occurrences, list) or not occurrences:
        raise CatalogMergeError(f"{context}.occurrences must be a non-empty list")
    occurrence_keys: set[tuple[str, int, int, str, int, int]] = set()
    for index, occurrence in enumerate(occurrences):
        key = _validate_occurrence(
            occurrence,
            context=f"{context}.occurrences[{index}]",
            sheet_names=sheet_names,
        )
        if key in occurrence_keys:
            raise CatalogMergeError(f"{context} contains a duplicate occurrence")
        occurrence_keys.add(key)
    return value


def _validate_member_asset_references(
    sheets: list[Any],
    *,
    context: str,
    asset_shas: set[str],
) -> None:
    for sheet_index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict) or not isinstance(sheet.get("sections"), list):
            raise CatalogMergeError(
                f"{context}.sheets[{sheet_index}] has invalid sections"
            )
        for section_index, section in enumerate(sheet["sections"]):
            if not isinstance(section, dict) or not isinstance(
                section.get("teams"), list
            ):
                raise CatalogMergeError(
                    f"{context}.sheets[{sheet_index}].sections[{section_index}] "
                    "has invalid teams"
                )
            for team_index, team in enumerate(section["teams"]):
                if not isinstance(team, dict) or not isinstance(
                    team.get("formation"), list
                ):
                    raise CatalogMergeError(
                        f"{context}.sheets[{sheet_index}].sections[{section_index}]."
                        f"teams[{team_index}] has an invalid formation"
                    )
                for member_index, member in enumerate(team["formation"]):
                    field = (
                        f"{context}.sheets[{sheet_index}].sections[{section_index}]."
                        f"teams[{team_index}].formation[{member_index}].image_sha256"
                    )
                    if not isinstance(member, dict):
                        raise CatalogMergeError(f"{field} has an invalid member")
                    digest = _require_sha256(member.get("image_sha256"), field=field)
                    if digest not in asset_shas:
                        raise CatalogMergeError(
                            f"{field} references an undeclared asset digest"
                        )


def validate_staging_catalog(
    document: Mapping[str, Any],
    *,
    context: str = "staging catalog",
) -> _ValidatedStaging:
    if not isinstance(document, dict):
        raise CatalogMergeError(f"{context} must be a JSON object")
    if set(document) != _TOP_LEVEL_KEYS:
        raise CatalogMergeError(f"{context} has an unsupported top-level shape")
    if document["schema_version"] != SCHEMA_VERSION:
        raise CatalogMergeError(
            f"{context} schema must be {SCHEMA_VERSION!r}; "
            f"got {document['schema_version']!r}"
        )
    if not isinstance(document["parser"], dict):
        raise CatalogMergeError(f"{context}.parser must be an object")
    if not isinstance(document["summary"], dict):
        raise CatalogMergeError(f"{context}.summary must be an object")
    if not isinstance(document["diagnostics"], list) or not all(
        isinstance(item, dict) for item in document["diagnostics"]
    ):
        raise CatalogMergeError(f"{context}.diagnostics must be a list of objects")

    source = _validate_source_workbook(document["source_workbook"], context=context)
    sheets = document["sheets"]
    if not isinstance(sheets, list) or not sheets:
        raise CatalogMergeError(f"{context}.sheets must be a non-empty list")
    sheet_names: set[str] = set()
    for index, sheet in enumerate(sheets):
        if not isinstance(sheet, dict) or not isinstance(sheet.get("sheet_name"), str):
            raise CatalogMergeError(f"{context}.sheets[{index}] has no sheet_name")
        name = sheet["sheet_name"]
        if name in sheet_names:
            raise CatalogMergeError(f"{context} contains duplicate sheet {name!r}")
        sheet_names.add(name)

    assets = document["assets"]
    if not isinstance(assets, list):
        raise CatalogMergeError(f"{context}.assets must be a list")
    assets_by_sha: dict[str, dict[str, Any]] = {}
    occupied_anchors: dict[tuple[str, int, int], str] = {}
    for index, value in enumerate(assets):
        asset = _validate_asset(
            value,
            context=f"{context}.assets[{index}]",
            sheet_names=sheet_names,
        )
        digest = asset["sha256"]
        if digest in assets_by_sha:
            raise CatalogMergeError(f"{context} contains duplicate asset {digest}")
        assets_by_sha[digest] = asset
        for occurrence in asset["occurrences"]:
            anchor = (
                occurrence["sheet_name"],
                occurrence["row"],
                occurrence["column"],
            )
            previous = occupied_anchors.get(anchor)
            if previous is not None and previous != digest:
                raise CatalogMergeError(
                    f"{context} assigns one image anchor to multiple digests"
                )
            occupied_anchors[anchor] = digest

    _validate_member_asset_references(
        sheets,
        context=context,
        asset_shas=set(assets_by_sha),
    )
    canonical_document = copy.deepcopy(document)
    semantic_sha256 = hashlib.sha256(
        _canonical_json_bytes(canonical_document)
    ).hexdigest()
    return _ValidatedStaging(
        document=canonical_document,
        source_sha256=source["sha256"],
        semantic_sha256=semantic_sha256,
        assets_by_sha=copy.deepcopy(assets_by_sha),
    )


def load_staging_catalog(path: str | Path) -> _ValidatedStaging:
    source = Path(path)
    if source.suffix.lower() != ".json":
        raise CatalogMergeError(f"staging catalog must use a .json filename: {source}")
    if not source.is_file():
        raise CatalogMergeError(f"staging catalog is not a file: {source}")
    try:
        document = json.loads(
            source.read_text(encoding="utf-8"),
            object_pairs_hook=_object_without_duplicate_keys,
            parse_constant=_reject_json_constant,
        )
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise CatalogMergeError(f"cannot read staging catalog {source}") from exc
    return validate_staging_catalog(document, context=str(source))


def _source_entry(staging: _ValidatedStaging) -> dict[str, Any]:
    document = staging.document
    return {
        "diagnostics": copy.deepcopy(document["diagnostics"]),
        "parser": copy.deepcopy(document["parser"]),
        "sheets": copy.deepcopy(document["sheets"]),
        "source_workbook": copy.deepcopy(document["source_workbook"]),
        "staging_catalog": {
            "schema_version": document["schema_version"],
            "sha256": staging.semantic_sha256,
        },
        "summary": copy.deepcopy(document["summary"]),
    }


def _count_nested(sources: Iterable[dict[str, Any]]) -> tuple[int, int, int]:
    section_count = 0
    team_count = 0
    axis_count = 0
    for source in sources:
        for sheet in source["sheets"]:
            sections = sheet["sections"]
            section_count += len(sections)
            for section in sections:
                teams = section["teams"]
                team_count += len(teams)
                axis_count += sum(len(team.get("axes", [])) for team in teams)
    return section_count, team_count, axis_count


def merge_staging_catalogs(
    documents: Iterable[Mapping[str, Any]],
) -> LocalCatalog:
    staged = [
        validate_staging_catalog(document, context=f"staging catalog {index + 1}")
        for index, document in enumerate(documents)
    ]
    if len(staged) < 2:
        raise CatalogMergeError("at least two staging catalogs are required")

    sources_by_sha: dict[str, _ValidatedStaging] = {}
    for item in staged:
        previous = sources_by_sha.get(item.source_sha256)
        if previous is not None:
            if previous.semantic_sha256 != item.semantic_sha256:
                raise CatalogMergeError(
                    "conflicting staging catalogs claim source workbook digest "
                    f"{item.source_sha256}"
                )
            continue
        sources_by_sha[item.source_sha256] = item

    ordered_staging = [sources_by_sha[key] for key in sorted(sources_by_sha)]
    sources = tuple(_source_entry(item) for item in ordered_staging)

    merged_assets: dict[str, dict[str, Any]] = {}
    occurrence_keys: dict[str, set[tuple[Any, ...]]] = {}
    for item in ordered_staging:
        for digest, staged_asset in sorted(item.assets_by_sha.items()):
            metadata = {key: staged_asset[key] for key in _ASSET_METADATA_KEYS}
            merged = merged_assets.get(digest)
            if merged is None:
                merged = {
                    **metadata,
                    "filename": f"{digest}.{staged_asset['file_extension']}",
                    "occurrences": [],
                    "sha256": digest,
                }
                merged_assets[digest] = merged
                occurrence_keys[digest] = set()
            elif any(merged[key] != metadata[key] for key in _ASSET_METADATA_KEYS):
                raise CatalogMergeError(
                    f"conflicting metadata for shared asset digest {digest}"
                )

            for occurrence in staged_asset["occurrences"]:
                qualified = {
                    **copy.deepcopy(occurrence),
                    "source_workbook_sha256": item.source_sha256,
                }
                key = (
                    qualified["source_workbook_sha256"],
                    qualified["sheet_name"],
                    qualified["row"],
                    qualified["column"],
                    qualified["anchor_cell"],
                    qualified["width_px"],
                    qualified["height_px"],
                )
                if key not in occurrence_keys[digest]:
                    merged["occurrences"].append(qualified)
                    occurrence_keys[digest].add(key)

    assets: list[dict[str, Any]] = []
    for digest in sorted(merged_assets):
        asset = merged_assets[digest]
        asset["occurrences"].sort(
            key=lambda item: (
                item["source_workbook_sha256"],
                item["sheet_name"],
                item["row"],
                item["column"],
                item["anchor_cell"],
                item["width_px"],
                item["height_px"],
            )
        )
        assets.append(asset)

    section_count, team_count, axis_count = _count_nested(sources)
    summary = {
        "asset_count": len(assets),
        "asset_occurrence_count": sum(
            len(asset["occurrences"]) for asset in assets
        ),
        "axis_count": axis_count,
        "diagnostic_count": sum(len(source["diagnostics"]) for source in sources),
        "section_count": section_count,
        "sheet_count": sum(len(source["sheets"]) for source in sources),
        "source_workbook_count": len(sources),
        "team_count": team_count,
    }
    return LocalCatalog(
        schema_version=LOCAL_CATALOG_SCHEMA_VERSION,
        builder={
            "name": CATALOG_BUILDER_NAME,
            "version": CATALOG_BUILDER_VERSION,
        },
        sources=sources,
        assets=tuple(assets),
        summary=summary,
    )


def merge_staging_catalog_files(paths: Iterable[str | Path]) -> LocalCatalog:
    resolved = [Path(path).resolve(strict=False) for path in paths]
    if len(set(resolved)) != len(resolved):
        raise CatalogMergeError("staging catalog input paths must be distinct")
    # Sorting makes both validation and successful output independent of CLI
    # input order, including on case-insensitive Windows filesystems.
    ordered = sorted(resolved, key=lambda path: str(path).casefold())
    documents = [load_staging_catalog(path).document for path in ordered]
    return merge_staging_catalogs(documents)


def _require_runtime_directory(path: str | Path) -> Path:
    destination = Path(path).resolve(strict=False)
    if ".runtime" not in {part.casefold() for part in destination.parts}:
        raise CatalogMergeError(
            "embedded media output must be inside a directory named .runtime"
        )
    if destination.exists() and not destination.is_dir():
        raise CatalogMergeError(
            f"embedded media output is not a directory: {destination}"
        )
    return destination


def export_staging_media(
    *,
    workbook_path: str | Path,
    staging_catalog_path: str | Path,
    output_directory: str | Path,
) -> dict[str, Any]:
    """Verify one workbook/catalog pair and export its exact embedded bytes."""

    workbook_file = Path(workbook_path).resolve(strict=False)
    if workbook_file.suffix.lower() != ".xlsx" or not workbook_file.is_file():
        raise CatalogMergeError(
            f"media source must be an existing .xlsx file: {workbook_file}"
        )
    staging = load_staging_catalog(staging_catalog_path)
    destination = _require_runtime_directory(output_directory)

    try:
        raw_bytes = workbook_file.read_bytes()
    except OSError as exc:
        raise CatalogMergeError(f"cannot read source workbook {workbook_file}") from exc
    source = staging.document["source_workbook"]
    actual_sha256 = hashlib.sha256(raw_bytes).hexdigest()
    if (
        actual_sha256 != source["sha256"]
        or len(raw_bytes) != source["byte_length"]
    ):
        raise CatalogMergeError(
            "source workbook bytes do not match staging catalog provenance"
        )

    workbook = load_workbook(BytesIO(raw_bytes), data_only=False, read_only=False)
    try:
        sheet_names = [sheet["sheet_name"] for sheet in staging.document["sheets"]]
        missing = [name for name in sheet_names if name not in workbook.sheetnames]
        if missing:
            raise CatalogMergeError(
                "source workbook is missing staged sheets: " + ", ".join(missing)
            )
        try:
            media = MediaCatalog.from_workbook(workbook, sheet_names)
        except WorkbookLayoutError as exc:
            raise CatalogMergeError(str(exc)) from exc
    finally:
        workbook.close()

    actual_assets = {
        asset.sha256: json.loads(json.dumps(asdict(asset)))
        for asset in media.assets()
    }
    if set(actual_assets) != set(staging.assets_by_sha):
        raise CatalogMergeError(
            "embedded media digests do not match the staging catalog"
        )
    for digest in sorted(actual_assets):
        if actual_assets[digest] != staging.assets_by_sha[digest]:
            raise CatalogMergeError(
                f"embedded media metadata conflicts with staging asset {digest}"
            )

    try:
        counts = media.export_exact_bytes(destination)
    except WorkbookLayoutError as exc:
        raise CatalogMergeError(str(exc)) from exc
    return {
        **counts,
        "source_workbook_sha256": actual_sha256,
    }
