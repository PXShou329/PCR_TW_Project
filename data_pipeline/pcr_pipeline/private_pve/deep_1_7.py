from __future__ import annotations

import hashlib
import re
import unicodedata
from dataclasses import replace
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from .media import MediaCatalog
from .models import (
    PARSER_NAME,
    PARSER_VERSION,
    SCHEMA_VERSION,
    AxisExtract,
    EraMarker,
    MemberExtract,
    SectionExtract,
    SheetExtract,
    StageRef,
    TeamExtract,
    WorkbookExtract,
    WorkbookLayoutError,
)
from .operations import merge_parsed_operations, parse_operation
from .sources import extract_source_refs


DEEP_SHEETS = (
    ("紅焔の深域(火屬性)", "FIRE"),
    ("蒼波の深域(水屬性)", "WATER"),
    ("翠嵐の深域(風屬性)", "WIND"),
    ("珀天の深域(光屬性)", "LIGHT"),
    ("紫冥の深域(闇屬性)", "DARK"),
)
ANEMONE_SHEET = "安涅默涅道中通關組合"
REQUIRED_SHEETS = tuple(name for name, _element in DEEP_SHEETS) + (ANEMONE_SHEET,)

_PANELS = (
    ("L", 1, 5, 6, 7, 9),
    ("M", 10, 14, 15, 16, 18),
    ("R", 19, 23, 24, 25, 27),
)
_ANEMONE_PANELS = (
    ("L", 1, 2, 6, 7, 8, 9),
    ("M", 2, 11, 15, 16, 17, 18),
    ("R", 3, 20, 24, 25, 26, 27),
)
_ANEMONE_BLOCKS = (
    ("FIRE", 3, 11, "FFEA9999"),
    ("WATER", 13, 21, "FFA4C2F4"),
    ("WIND", 23, 31, "FFB6D7A8"),
    ("LIGHT", 33, 41, "FFFFE599"),
    ("DARK", 43, 51, "FFB4A7D6"),
)
_STAGE_PATTERN = re.compile(r"(?<!\d)([1-7])[-－](10|[1-9])(?!\d)")


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _merged_bounds(worksheet, row: int, column: int) -> tuple[int, int, int, int]:
    for merged in worksheet.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= column <= merged.max_col
        ):
            return merged.min_row, merged.min_col, merged.max_row, merged.max_col
    return row, column, row, column


def _range_label(bounds: tuple[int, int, int, int]) -> str:
    min_row, min_col, max_row, max_col = bounds
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _merged_anchor_cell(worksheet, row: int, column: int):
    min_row, min_col, _max_row, _max_col = _merged_bounds(worksheet, row, column)
    return worksheet.cell(min_row, min_col)


def _parse_stage_label(value: Any) -> tuple[int, int] | None:
    if not isinstance(value, str):
        return None
    normalized = unicodedata.normalize("NFKC", value)
    normalized = normalized.replace(" ", "").replace("\u3000", "")
    match = _STAGE_PATTERN.search(normalized)
    if match is None:
        return None
    return int(match.group(1)), int(match.group(2))


def _formation(
    media: MediaCatalog,
    *,
    sheet_name: str,
    row: int,
    first_column: int,
) -> tuple[MemberExtract, ...] | None:
    digests = [
        media.image_sha(sheet_name, row, column)
        for column in range(first_column, first_column + 5)
    ]
    present = sum(digest is not None for digest in digests)
    if present == 0:
        return None
    if present != 5:
        start = f"{get_column_letter(first_column)}{row}"
        raise WorkbookLayoutError(
            f"{sheet_name}!{start} roster has {present} images; expected exactly 5"
        )
    return tuple(
        MemberExtract(
            display_position=slot,
            anchor_cell=f"{get_column_letter(first_column + slot - 1)}{row}",
            image_sha256=digest,
        )
        for slot, digest in enumerate(digests, start=1)
        if digest is not None
    )


def _formation_signature(team: TeamExtract) -> tuple[str, ...]:
    return tuple(member.image_sha256 for member in team.formation)


def _team_source_flags(axes: tuple[AxisExtract, ...]) -> tuple[str, ...]:
    with_sources = sum(bool(axis.source_refs) for axis in axes)
    if with_sources == 0:
        return ("NO_SOURCE_REFERENCE",)
    if with_sources != len(axes):
        return ("SOME_AXES_NO_SOURCE_REFERENCE",)
    return ()


def _merge_adjacent_team_axes(
    ordered: list[tuple[int, int, TeamExtract]],
) -> list[tuple[int, int, TeamExtract]]:
    """Merge adjacent duplicate portrait rows into one team with multiple axes."""

    groups: list[tuple[int, int, int, TeamExtract]] = []
    for panel_index, row, team in ordered:
        if groups:
            previous_panel, first_row, last_row, previous = groups[-1]
            if (
                previous_panel == panel_index
                and row == last_row + 1
                and _formation_signature(previous) == _formation_signature(team)
            ):
                axes = previous.axes + team.axes
                refs = _dedupe_refs(previous.source_refs + team.source_refs)
                first_slot = _PANELS[panel_index][1]
                source_col = _PANELS[panel_index][5]
                groups[-1] = (
                    previous_panel,
                    first_row,
                    row,
                    replace(
                        previous,
                        source_range=(
                            f"{get_column_letter(first_slot)}{first_row}:"
                            f"{get_column_letter(source_col)}{row}"
                        ),
                        axes=axes,
                        notes_raw=(
                            previous.notes_raw
                            if previous.notes_raw == team.notes_raw
                            else None
                        ),
                        source_refs=refs,
                        flags=_team_source_flags(axes),
                    ),
                )
                continue
        groups.append((panel_index, row, row, team))
    return [(panel, first_row, team) for panel, first_row, _last_row, team in groups]


def _parse_deep_sheet(worksheet, media: MediaCatalog, *, element: str, strict: bool):
    current: dict[str, tuple[int, int] | None] = {panel[0]: None for panel in _PANELS}
    labels: dict[tuple[int, int], tuple[str, str]] = {}
    drafts: dict[tuple[int, int], list[tuple[int, int, TeamExtract]]] = {}

    for row in range(1, worksheet.max_row + 1):
        for panel_index, (panel, first_slot, _last_slot, _op, _notes, _source) in enumerate(
            _PANELS
        ):
            cell = worksheet.cell(row, first_slot)
            parsed = _parse_stage_label(cell.value)
            if parsed is None:
                continue
            bounds = _merged_bounds(worksheet, row, first_slot)
            valid_header = (
                bounds[0] == row
                and bounds[2] == row
                and bounds[1] == first_slot
                and bounds[3] in {9, 18, 27}
                and bounds[3] >= first_slot + 8
            )
            if not valid_header:
                continue
            label = _cell_text(cell.value) or f"{parsed[0]}-{parsed[1]}"
            source_range = _range_label(bounds)
            previous = labels.get(parsed)
            if previous is not None and previous != (label, source_range):
                raise WorkbookLayoutError(
                    f"{worksheet.title} contains duplicate header for {parsed[0]}-{parsed[1]}"
                )
            labels[parsed] = (label, source_range)
            for target_panel, target_first_slot, *_rest in _PANELS:
                if bounds[1] <= target_first_slot <= bounds[3]:
                    current[target_panel] = parsed

        for panel_index, (panel, first_slot, last_slot, op_col, notes_col, source_col) in enumerate(
            _PANELS
        ):
            formation = _formation(
                media,
                sheet_name=worksheet.title,
                row=row,
                first_column=first_slot,
            )
            if formation is None:
                continue
            stage = current[panel]
            if stage is None:
                raise WorkbookLayoutError(
                    f"{worksheet.title}!{get_column_letter(first_slot)}{row} has no stage header"
                )
            operation_cell = _merged_anchor_cell(worksheet, row, op_col)
            operation_raw = _cell_text(operation_cell.value)
            if operation_raw is None or not operation_raw.strip():
                raise WorkbookLayoutError(
                    f"{worksheet.title}!{operation_cell.coordinate} has no operation"
                )
            notes_cell = _merged_anchor_cell(worksheet, row, notes_col)
            notes_raw = _cell_text(notes_cell.value)
            requested_source_cell = worksheet.cell(row, source_col)
            source_cell = _merged_anchor_cell(worksheet, row, source_col)
            source_refs = extract_source_refs(
                source_cell,
                sheet_name=worksheet.title,
                applies_to_cell=(
                    requested_source_cell.coordinate
                    if requested_source_cell.coordinate != source_cell.coordinate
                    else None
                ),
            )
            team_id = (
                f"deep:{element.lower()}:{stage[0]:02d}-{stage[1]:02d}:"
                f"{panel.lower()}:{row:03d}"
            )
            axis = AxisExtract(
                axis_source_id=f"{team_id}:axis:{operation_cell.coordinate.lower()}",
                source_cell=operation_cell.coordinate,
                operation_raw=operation_raw,
                operation=merge_parsed_operations(
                    parse_operation(
                        operation_raw,
                        origin_field="OPERATION_TEXT",
                        origin_cell=operation_cell.coordinate,
                    ),
                    parse_operation(
                        notes_raw or "",
                        origin_field="NOTES_TEXT",
                        origin_cell=notes_cell.coordinate,
                    ),
                ),
                notes_raw=notes_raw,
                source_refs=source_refs,
            )
            flags = () if source_refs else ("NO_SOURCE_REFERENCE",)
            team = TeamExtract(
                team_source_id=team_id,
                display_order=0,
                source_range=(
                    f"{get_column_letter(first_slot)}{row}:"
                    f"{get_column_letter(source_col)}{row}"
                ),
                formation=formation,
                axes=(axis,),
                notes_raw=notes_raw,
                source_refs=source_refs,
                flags=flags,
            )
            drafts.setdefault(stage, []).append((panel_index, row, team))

    if strict:
        expected = {(area, stage) for area in range(1, 8) for stage in range(1, 11)}
        if set(drafts) != expected or set(labels) != expected:
            missing = sorted(expected - set(drafts))
            extra = sorted(set(drafts) - expected)
            raise WorkbookLayoutError(
                f"{worksheet.title} stage closure mismatch; missing={missing}, extra={extra}"
            )

    sections: list[SectionExtract] = []
    for area, stage in sorted(drafts):
        ordered = _merge_adjacent_team_axes(
            sorted(drafts[(area, stage)], key=lambda item: (item[0], item[1]))
        )
        teams = tuple(
            replace(team, display_order=index)
            for index, (_panel, _row, team) in enumerate(ordered, start=1)
        )
        label, source_range = labels.get(
            (area, stage), (f"{area}-{stage}", teams[0].source_range)
        )
        sections.append(
            SectionExtract(
                section_id=f"deep:{element.lower()}:{area:02d}-{stage:02d}",
                mode="DEEP",
                element=element,
                stage_refs=(StageRef(kind="DEEP", area=area, stage=stage),),
                label_raw=label,
                source_range=source_range,
                era=EraMarker(kind="UNKNOWN"),
                teams=teams,
            )
        )
    return SheetExtract(
        sheet_name=worksheet.title,
        sheet_state=worksheet.sheet_state,
        layout_id="deep-1-7-three-panel/v1",
        collection="DEEP_1_7_WORKBOOK",
        element=element,
        source_refs=(),
        sections=tuple(sections),
    )


def _anemone_stage(value: Any, *, epoch) -> tuple[int, int] | None:
    if isinstance(value, datetime):
        return value.month, value.day
    if isinstance(value, date):
        return value.month, value.day
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch=epoch)
        return converted.month, converted.day
    return _parse_stage_label(value)


def _dedupe_refs(refs):
    result = []
    seen = set()
    for ref in refs:
        key = (
            ref.kind,
            ref.url,
            ref.alias_id,
            ref.origin_cell,
            ref.applies_to_cell,
        )
        if key not in seen:
            result.append(ref)
            seen.add(key)
    return tuple(result)


def _parse_anemone_sheet(worksheet, media: MediaCatalog, *, epoch, strict: bool):
    sections: list[SectionExtract] = []
    for element, first_row, last_row, expected_fill in _ANEMONE_BLOCKS:
        element_sections: list[SectionExtract] = []
        for row in range(first_row, last_row + 1):
            expected_stage = row - first_row + 1
            for panel, area, first_slot, _last_slot, op_col, notes_col, source_col in _ANEMONE_PANELS:
                formation = _formation(
                    media,
                    sheet_name=worksheet.title,
                    row=row,
                    first_column=first_slot,
                )
                stage_cell = worksheet.cell(row, first_slot - 1)
                if formation is None:
                    continue
                fill = stage_cell.fill.fgColor.rgb
                if fill != expected_fill:
                    raise WorkbookLayoutError(
                        f"{worksheet.title}!{stage_cell.coordinate} fill {fill!r} "
                        f"does not match {element} block {expected_fill}"
                    )
                parsed = _anemone_stage(stage_cell.value, epoch=epoch)
                if parsed != (area, expected_stage):
                    raise WorkbookLayoutError(
                        f"{worksheet.title}!{stage_cell.coordinate} stage {parsed!r} "
                        f"does not match structural stage {(area, expected_stage)!r}"
                    )
                operation_cell = _merged_anchor_cell(worksheet, row, op_col)
                operation_raw = _cell_text(operation_cell.value)
                if operation_raw is None or not operation_raw.strip():
                    raise WorkbookLayoutError(
                        f"{worksheet.title}!{operation_cell.coordinate} has no operation"
                    )
                notes_cell = _merged_anchor_cell(worksheet, row, notes_col)
                notes_raw = _cell_text(notes_cell.value)
                requested_source_cell = worksheet.cell(row, source_col)
                source_cell = _merged_anchor_cell(worksheet, row, source_col)
                source_refs = extract_source_refs(
                    source_cell,
                    sheet_name=worksheet.title,
                    applies_to_cell=(
                        requested_source_cell.coordinate
                        if requested_source_cell.coordinate != source_cell.coordinate
                        else None
                    ),
                )
                team_id = (
                    f"anemone:{element.lower()}:{area:02d}-{expected_stage:02d}:"
                    f"{panel.lower()}:{row:03d}"
                )
                axis = AxisExtract(
                    axis_source_id=f"{team_id}:axis:{operation_cell.coordinate.lower()}",
                    source_cell=operation_cell.coordinate,
                    operation_raw=operation_raw,
                    operation=merge_parsed_operations(
                        parse_operation(
                            operation_raw,
                            origin_field="OPERATION_TEXT",
                            origin_cell=operation_cell.coordinate,
                        ),
                        parse_operation(
                            notes_raw or "",
                            origin_field="NOTES_TEXT",
                            origin_cell=notes_cell.coordinate,
                        ),
                    ),
                    notes_raw=notes_raw,
                    source_refs=source_refs,
                )
                team = TeamExtract(
                    team_source_id=team_id,
                    display_order=1,
                    source_range=(
                        f"{stage_cell.coordinate}:"
                        f"{get_column_letter(source_col)}{row}"
                    ),
                    formation=formation,
                    axes=(axis,),
                    notes_raw=notes_raw,
                    source_refs=source_refs,
                    flags=() if source_refs else ("NO_SOURCE_REFERENCE",),
                )
                element_sections.append(
                    SectionExtract(
                        section_id=(
                            f"anemone:{element.lower()}:{area:02d}-{expected_stage:02d}"
                        ),
                        mode="DEEP",
                        element=element,
                        stage_refs=(
                            StageRef(kind="DEEP", area=area, stage=expected_stage),
                        ),
                        # Google Sheets serialized these displayed `m-d` labels as
                        # dates.  The year is an import artifact, not part of the
                        # stage identity or the user-visible label.
                        label_raw=f"{area}-{expected_stage}",
                        source_range=stage_cell.coordinate,
                        era=EraMarker(kind="UNKNOWN"),
                        teams=(team,),
                    )
                )
        if strict and len(element_sections) != 27:
            raise WorkbookLayoutError(
                f"{worksheet.title} {element} block has {len(element_sections)} teams; expected 27"
            )
        sections.extend(element_sections)

    sheet_refs = []
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.hyperlink is not None:
                sheet_refs.extend(extract_source_refs(cell, sheet_name=worksheet.title))
    return SheetExtract(
        sheet_name=worksheet.title,
        sheet_state=worksheet.sheet_state,
        layout_id="anemone-road-clear-five-block/v1",
        collection="ANEMONE_ROAD_CLEAR",
        element=None,
        source_refs=_dedupe_refs(sheet_refs),
        sections=tuple(
            sorted(
                sections,
                key=lambda item: (
                    next(i for i, (_name, element) in enumerate(DEEP_SHEETS) if element == item.element),
                    item.stage_refs[0].area,
                    item.stage_refs[0].stage,
                ),
            )
        ),
    )


def extract_deep_1_7_workbook(
    path: str | Path,
    *,
    strict: bool = True,
) -> WorkbookExtract:
    input_path = Path(path)
    raw_bytes = input_path.read_bytes()
    # Parse the exact bytes that were hashed.  Loading the path a second time
    # would permit a concurrently replaced workbook to produce a catalog whose
    # declared source digest belongs to different content.
    workbook = load_workbook(BytesIO(raw_bytes), read_only=False, data_only=False)
    try:
        missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
        if missing:
            raise WorkbookLayoutError(
                "workbook is not the deep 1-7 source; missing sheets: "
                + ", ".join(missing)
            )
        media = MediaCatalog.from_workbook(workbook, REQUIRED_SHEETS)
        sheets = [
            _parse_deep_sheet(
                workbook[sheet_name], media, element=element, strict=strict
            )
            for sheet_name, element in DEEP_SHEETS
        ]
        sheets.append(
            _parse_anemone_sheet(
                workbook[ANEMONE_SHEET],
                media,
                epoch=workbook.epoch,
                strict=strict,
            )
        )
        assets = media.assets()
    finally:
        workbook.close()

    section_count = sum(len(sheet.sections) for sheet in sheets)
    team_count = sum(
        len(section.teams) for sheet in sheets for section in sheet.sections
    )
    axis_count = sum(
        len(team.axes)
        for sheet in sheets
        for section in sheet.sections
        for team in section.teams
    )
    source_ref_count = sum(len(sheet.source_refs) for sheet in sheets) + sum(
        len(team.source_refs)
        for sheet in sheets
        for section in sheet.sections
        for team in section.teams
    )
    diagnostics = tuple(
        {
            "code": "SHEET_OUT_OF_SCOPE",
            "sheet_name": sheet_name,
            "sheet_state": workbook_state,
        }
        for sheet_name, workbook_state in sorted(
            (
                (sheet_name, "UNKNOWN")
                for sheet_name in set(workbook.sheetnames) - set(REQUIRED_SHEETS)
            ),
            key=lambda item: item[0],
        )
    )
    summary = {
        "asset_count": len(assets),
        "asset_occurrence_count": sum(len(asset.occurrences) for asset in assets),
        "axis_count": axis_count,
        "section_count": section_count,
        "sheet_count": len(sheets),
        "source_ref_count": source_ref_count,
        "team_count": team_count,
        "team_counts_by_sheet": {
            sheet.sheet_name: sum(len(section.teams) for section in sheet.sections)
            for sheet in sheets
        },
        "unmapped_member_count": team_count * 5,
    }
    return WorkbookExtract(
        schema_version=SCHEMA_VERSION,
        parser={
            "layout_ids": [
                "deep-1-7-three-panel/v1",
                "anemone-road-clear-five-block/v1",
            ],
            "name": PARSER_NAME,
            "version": PARSER_VERSION,
        },
        source_workbook={
            "byte_length": len(raw_bytes),
            "filename": input_path.name,
            "sha256": hashlib.sha256(raw_bytes).hexdigest(),
        },
        sheets=tuple(sheets),
        assets=assets,
        summary=summary,
        diagnostics=diagnostics,
    )
