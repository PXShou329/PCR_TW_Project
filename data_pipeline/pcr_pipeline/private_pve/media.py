from __future__ import annotations

import hashlib
import os
import tempfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl.utils import get_column_letter

from .models import AssetExtract, AssetOccurrence, WorkbookLayoutError


def _media_type(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png", "png"
    if data.startswith(b"\xff\xd8\xff"):
        return "image/jpeg", "jpg"
    if data.startswith((b"GIF87a", b"GIF89a")):
        return "image/gif", "gif"
    if len(data) >= 12 and data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return "image/webp", "webp"
    raise WorkbookLayoutError("unsupported embedded image format")


def _anchor(image) -> tuple[int, int]:
    anchor = image.anchor
    if isinstance(anchor, str):
        from openpyxl.utils.cell import coordinate_to_tuple

        return coordinate_to_tuple(anchor)
    marker = getattr(anchor, "_from", None)
    if marker is None:
        raise WorkbookLayoutError("embedded image has no supported anchor marker")
    return marker.row + 1, marker.col + 1


@dataclass
class _MutableAsset:
    sha256: str
    byte_length: int
    mime_type: str
    file_extension: str
    width_px: int
    height_px: int
    source_bytes: bytes
    occurrences: list[AssetOccurrence]


def _verify_existing_asset(path: Path, expected: bytes, digest: str) -> None:
    try:
        actual = path.read_bytes()
    except OSError as exc:
        raise WorkbookLayoutError(
            f"cannot verify existing exported asset {path}"
        ) from exc
    if actual != expected or hashlib.sha256(actual).hexdigest() != digest:
        raise WorkbookLayoutError(
            f"existing exported asset does not match {digest}: {path}"
        )


def _write_asset_atomic(path: Path, payload: bytes, digest: str) -> bool:
    """Create one content-addressed asset without ever replacing an existing file."""

    if path.exists():
        _verify_existing_asset(path, payload, digest)
        return False

    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary_path = Path(handle.name)
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())

        try:
            # A same-directory hard link is an atomic create-if-absent operation.
            # Unlike os.replace(), it cannot overwrite a file created by another
            # process between the initial existence check and this operation.
            os.link(temporary_path, path)
        except FileExistsError:
            _verify_existing_asset(path, payload, digest)
            return False
        return True
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


class MediaCatalog:
    """Exact-byte image inventory keyed by sheet cell and SHA-256."""

    def __init__(self) -> None:
        self._by_cell: dict[tuple[str, int, int], list[str]] = defaultdict(list)
        self._assets: dict[str, _MutableAsset] = {}

    @classmethod
    def from_workbook(cls, workbook, sheet_names: Iterable[str]) -> "MediaCatalog":
        catalog = cls()
        for sheet_name in sheet_names:
            worksheet = workbook[sheet_name]
            for image in worksheet._images:
                source_format = str(getattr(image, "format", "")).lower()
                if source_format not in {"gif", "jpeg", "jpg", "png"}:
                    raise WorkbookLayoutError(
                        "unsupported embedded source image format "
                        f"{source_format!r}; refusing an implicit Pillow conversion"
                    )
                row, column = _anchor(image)
                data = image._data()
                digest = hashlib.sha256(data).hexdigest()
                mime_type, extension = _media_type(data)
                source_extension = "jpg" if source_format == "jpeg" else source_format
                if source_extension != extension:
                    raise WorkbookLayoutError(
                        "embedded image format metadata does not match its source bytes; "
                        "refusing an implicit conversion"
                    )
                occurrence = AssetOccurrence(
                    sheet_name=sheet_name,
                    anchor_cell=f"{get_column_letter(column)}{row}",
                    row=row,
                    column=column,
                    width_px=int(round(image.width)),
                    height_px=int(round(image.height)),
                )
                catalog._by_cell[(sheet_name, row, column)].append(digest)
                existing = catalog._assets.get(digest)
                if existing is None:
                    catalog._assets[digest] = _MutableAsset(
                        sha256=digest,
                        byte_length=len(data),
                        mime_type=mime_type,
                        file_extension=extension,
                        width_px=occurrence.width_px,
                        height_px=occurrence.height_px,
                        source_bytes=data,
                        occurrences=[occurrence],
                    )
                else:
                    if (
                        existing.byte_length != len(data)
                        or existing.mime_type != mime_type
                        or existing.file_extension != extension
                        or existing.source_bytes != data
                    ):
                        raise WorkbookLayoutError(
                            f"embedded media hash collision for {digest}"
                        )
                    existing.occurrences.append(occurrence)
        return catalog

    def image_sha(self, sheet_name: str, row: int, column: int) -> str | None:
        matches = self._by_cell.get((sheet_name, row, column), [])
        if len(matches) > 1:
            cell = f"{get_column_letter(column)}{row}"
            raise WorkbookLayoutError(
                f"{sheet_name}!{cell} contains more than one embedded image"
            )
        return matches[0] if matches else None

    def occurrence_keys(self) -> set[tuple[str, int, int, str]]:
        return {
            (sheet, row, column, digest)
            for (sheet, row, column), digests in self._by_cell.items()
            for digest in digests
        }

    def export_exact_bytes(self, output_directory: str | Path) -> dict[str, int]:
        """Export exact embedded bytes under their SHA-256 filenames.

        Existing files are never replaced.  A byte-for-byte match is treated as
        an idempotent verification; any mismatch fails closed.
        """

        destination = Path(output_directory)
        destination.mkdir(parents=True, exist_ok=True)
        if not destination.is_dir():
            raise WorkbookLayoutError(
                f"media export destination is not a directory: {destination}"
            )

        written_count = 0
        verified_existing_count = 0
        for digest in sorted(self._assets):
            asset = self._assets[digest]
            if hashlib.sha256(asset.source_bytes).hexdigest() != digest:
                raise WorkbookLayoutError(
                    f"embedded media bytes no longer match catalog digest {digest}"
                )
            path = destination / f"{digest}.{asset.file_extension}"
            if _write_asset_atomic(path, asset.source_bytes, digest):
                written_count += 1
            else:
                verified_existing_count += 1
        return {
            "asset_count": len(self._assets),
            "verified_existing_count": verified_existing_count,
            "written_count": written_count,
        }

    def assets(self) -> tuple[AssetExtract, ...]:
        result: list[AssetExtract] = []
        for digest in sorted(self._assets):
            asset = self._assets[digest]
            occurrences = tuple(
                sorted(
                    asset.occurrences,
                    key=lambda item: (
                        item.sheet_name,
                        item.row,
                        item.column,
                        item.anchor_cell,
                    ),
                )
            )
            result.append(
                AssetExtract(
                    sha256=asset.sha256,
                    byte_length=asset.byte_length,
                    mime_type=asset.mime_type,
                    file_extension=asset.file_extension,
                    width_px=asset.width_px,
                    height_px=asset.height_px,
                    occurrences=occurrences,
                )
            )
        return tuple(result)
