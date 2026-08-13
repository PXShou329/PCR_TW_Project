from __future__ import annotations

import hashlib
import re
import unicodedata
from collections import defaultdict
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from .media import MediaCatalog
from .models import (
    PARSER_NAME,
    PARSER_VERSION,
    SCHEMA_VERSION,
    AxisExtract,
    EraMarker,
    MemberExtract,
    SectionExtract,
    SheetExtract,
    StageRef,
    TeamExtract,
    WorkbookExtract,
    WorkbookLayoutError,
)
from .operations import merge_parsed_operations, parse_operation
from .sources import extract_source_refs


DEEP_8_10_SHEETS = (
    ("火8", "FIRE", 8),
    ("水8", "WATER", 8),
    ("風8", "WIND", 8),
    ("光8", "LIGHT", 8),
    ("闇8", "DARK", 8),
    ("火9", "FIRE", 9),
    ("水9", "WATER", 9),
    ("風9", "WIND", 9),
    ("光9", "LIGHT", 9),
    ("闇9", "DARK", 9),
    ("火10", "FIRE", 10),
    ("水10", "WATER", 10),
    ("風10", "WIND", 10),
    ("光10", "LIGHT", 10),
    ("闇10", "DARK", 10),
)

REMEMBRANCE_SHEETS = (
    ("追憶霸瞳6-10", "HATOU", ((6, 7, 8), (9,), (10,))),
    ("追憶贊恩6-10", "ZANE", ((6, 7, 8), (9,), (10,))),
    (
        "追憶彌勒1-10",
        "MAITREYA",
        ((1, 2, 3), (4,), (5,), (6, 7, 8), (9,), (10,)),
    ),
    ("追憶阿剌克涅1-5", "ARACHNE", ((1, 2, 3), (4,), (5,))),
)

LUNA_SHEET = "露娜塔頂層EX"
TARGET_SHEETS = (
    tuple(name for name, _element, _area in DEEP_8_10_SHEETS)
    + tuple(name for name, _boss, _groups in REMEMBRANCE_SHEETS)
    + (LUNA_SHEET,)
)

_DEEP_STAGE_PATTERN = re.compile(
    r"(?<!\d)(8|9|10)\s*[-－]\s*(10|[1-9])(?!\d)"
)
_FLOOR_PATTERN = re.compile(
    r"(?<!\d)(\d{1,2})(?:\s*[-－~]\s*(\d{1,2}))?\s*層"
)
_HYPERLINK_PREFIX = re.compile(r"^\s*=\s*HYPERLINK\b", re.IGNORECASE)

_DEEP_PANELS = (
    ("L", 2, 7, 8, 9),
    ("R", 12, 17, 18, 19),
)

# The divider column belongs to the left panel when it contains a linked
# alternate-operation label.  The current workbook uses this at 火8!J28 for
# the left roster's 全SET alternative video.
_DEEP_LEFT_EXTRA_SOURCE_COLUMN = 10

# These portrait-sized runs are audited workbook illustrations/examples, not
# playable formations.  Any new unconsumed run of five or more portraits is a
# layout change and must fail closed instead of silently dropping a team.
_AUDITED_UNCONSUMED_PORTRAIT_RUNS = frozenset(
    {
        *((f"{element}10", tuple(f"{column}3" for column in "ABCDEF")) for element in "火水風光闇"),
        *((f"{element}10", tuple(f"{column}7" for column in "BCDEF")) for element in "火水風光闇"),
        ("追憶阿剌克涅1-5", tuple(f"{column}8" for column in "BCDEF")),
    }
)

# The detailed Luna Tower area is prose-oriented rather than a rectangular
# table.  These anchors are audited source coordinates, not inferred ranges.
_LUNA_DETAIL_BLOCKS = (
    ("current-party-1-method-1", 11, ("A12",), ("A10",)),
    ("current-party-1-method-2", 13, ("A14",), ()),
    ("current-party-2-method-1", 19, ("A20",), ("A18",)),
    ("current-party-2-method-2", 21, ("A22",), ()),
    ("current-party-3", 25, ("A26",), ("F25",)),
    ("reference-party-1", 30, ("A31",), ("A29",)),
    ("reference-party-2", 36, ("A37",), ("A35",)),
    ("reference-party-3", 40, ("A41", "A43"), ("F40", "A42")),
)


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    return str(value).replace("\r\n", "\n").replace("\r", "\n")


def _merged_bounds(worksheet, row: int, column: int) -> tuple[int, int, int, int]:
    for merged in worksheet.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= column <= merged.max_col
        ):
            return merged.min_row, merged.min_col, merged.max_row, merged.max_col
    return row, column, row, column


def _merged_anchor_cell(worksheet, row: int, column: int):
    min_row, min_col, _max_row, _max_col = _merged_bounds(worksheet, row, column)
    return worksheet.cell(min_row, min_col)


def _range_label(bounds: tuple[int, int, int, int]) -> str:
    min_row, min_col, max_row, max_col = bounds
    first = f"{get_column_letter(min_col)}{min_row}"
    last = f"{get_column_letter(max_col)}{max_row}"
    return first if first == last else f"{first}:{last}"


def _last_content_row(worksheet) -> int:
    image_rows = []
    for image in worksheet._images:
        if isinstance(image.anchor, str):
            row, _column = coordinate_to_tuple(image.anchor)
        else:
            row = image.anchor._from.row + 1
        image_rows.append(row)
    return max([worksheet.max_row, *image_rows])


def _dedupe_refs(refs):
    result = []
    seen = set()
    for ref in refs:
        key = (
            ref.kind,
            ref.url,
            ref.alias_id,
            ref.origin_cell,
            ref.applies_to_cell,
        )
        if key not in seen:
            result.append(ref)
            seen.add(key)
    return tuple(result)


def _sheet_source_refs(worksheet):
    refs = []
    for row in worksheet.iter_rows():
        for cell in row:
            is_hyperlink_formula = (
                isinstance(cell.value, str)
                and _HYPERLINK_PREFIX.match(cell.value) is not None
            )
            if cell.hyperlink is not None or is_hyperlink_formula:
                refs.extend(extract_source_refs(cell, sheet_name=worksheet.title))
    return _dedupe_refs(refs)


def _note_cells(worksheet, row: int, first_column: int, last_column: int):
    result = []
    seen = set()
    for column in range(first_column, last_column + 1):
        cell = _merged_anchor_cell(worksheet, row, column)
        if cell.coordinate in seen:
            continue
        seen.add(cell.coordinate)
        result.append(cell)
    return tuple(result)


def _note_parts(worksheet, row: int, first_column: int, last_column: int):
    result = []
    for cell in _note_cells(worksheet, row, first_column, last_column):
        if cell.data_type == "f":
            continue
        raw = _cell_text(cell.value)
        if raw is not None and raw.strip():
            result.append((cell, raw))
    return tuple(result)


def _refs_for_cells(worksheet, cells, *, applies_to_cell: str | None = None):
    refs = []
    seen_cells = set()
    for cell in cells:
        if cell.coordinate in seen_cells:
            continue
        seen_cells.add(cell.coordinate)
        refs.extend(
            extract_source_refs(
                cell,
                sheet_name=worksheet.title,
                applies_to_cell=(
                    applies_to_cell
                    if applies_to_cell is not None
                    and cell.coordinate != applies_to_cell
                    else None
                ),
            )
        )
    return _dedupe_refs(refs)


def _make_axis(
    worksheet,
    *,
    axis_id: str,
    operation_cell,
    note_parts,
    extra_source_cells=(),
) -> AxisExtract:
    if operation_cell.data_type == "f":
        raise WorkbookLayoutError(
            f"{worksheet.title}!{operation_cell.coordinate} operation is a formula"
        )
    operation_raw = _cell_text(operation_cell.value)
    if operation_raw is None or not operation_raw.strip():
        raise WorkbookLayoutError(
            f"{worksheet.title}!{operation_cell.coordinate} has no operation"
        )
    parsed_fields = [
        parse_operation(
            operation_raw,
            origin_field="OPERATION_TEXT",
            origin_cell=operation_cell.coordinate,
        )
    ]
    for cell, raw in note_parts:
        parsed_fields.append(
            parse_operation(
                raw,
                origin_field="NOTES_TEXT",
                origin_cell=cell.coordinate,
            )
        )
    notes_raw = "\n".join(raw for _cell, raw in note_parts) or None
    source_cells = (operation_cell,) + tuple(cell for cell, _raw in note_parts)
    source_refs = _refs_for_cells(
        worksheet,
        source_cells + tuple(extra_source_cells),
        applies_to_cell=operation_cell.coordinate,
    )
    return AxisExtract(
        axis_source_id=axis_id,
        source_cell=operation_cell.coordinate,
        operation_raw=operation_raw,
        operation=merge_parsed_operations(*parsed_fields),
        notes_raw=notes_raw,
        source_refs=source_refs,
    )


def _formation_block(
    worksheet,
    media: MediaCatalog,
    *,
    row: int,
    first_column: int,
) -> tuple[tuple[MemberExtract, ...], int] | None:
    digests = [
        media.image_sha(worksheet.title, row, column)
        for column in range(first_column, first_column + 5)
    ]
    present = sum(digest is not None for digest in digests)
    if present == 0:
        return None
    if present != 5:
        raise WorkbookLayoutError(
            f"{worksheet.title}!{get_column_letter(first_column)}{row} "
            f"roster has {present} images; expected exactly 5"
        )
    bounds = [
        _merged_bounds(worksheet, row, column)
        for column in range(first_column, first_column + 5)
    ]
    if any(bound[0] != row for bound in bounds):
        raise WorkbookLayoutError(
            f"{worksheet.title}!{get_column_letter(first_column)}{row} "
            "image is not anchored at the top of its merged roster block"
        )
    bottom_rows = {bound[2] for bound in bounds}
    if len(bottom_rows) != 1:
        raise WorkbookLayoutError(
            f"{worksheet.title}!{get_column_letter(first_column)}{row} "
            "five roster cells have different vertical spans"
        )
    members = tuple(
        MemberExtract(
            display_position=index,
            anchor_cell=f"{get_column_letter(first_column + index - 1)}{row}",
            image_sha256=digest,
        )
        for index, digest in enumerate(digests, start=1)
        if digest is not None
    )
    return members, bottom_rows.pop()


def _axes_for_formation(
    worksheet,
    *,
    team_id: str,
    first_row: int,
    last_row: int,
    operation_column: int,
    notes_first_column: int,
    notes_last_column: int,
    extra_source_column: int | None = None,
) -> tuple[AxisExtract, ...]:
    axes = []
    seen_operations = set()
    for row in range(first_row, last_row + 1):
        operation_cell = _merged_anchor_cell(worksheet, row, operation_column)
        if operation_cell.coordinate in seen_operations:
            continue
        seen_operations.add(operation_cell.coordinate)
        if operation_cell.value is None:
            continue
        notes = _note_parts(
            worksheet,
            row,
            notes_first_column,
            notes_last_column,
        )
        note_cells = _note_cells(
            worksheet,
            row,
            notes_first_column,
            notes_last_column,
        )
        extra_source_cells = ()
        if extra_source_column is not None:
            extra_cell = _merged_anchor_cell(worksheet, row, extra_source_column)
            if extra_cell.value is not None or extra_cell.hyperlink is not None:
                extra_source_cells = (extra_cell,)
        axes.append(
            _make_axis(
                worksheet,
                axis_id=(
                    f"{team_id}:axis:{operation_cell.coordinate.lower()}"
                ),
                operation_cell=operation_cell,
                note_parts=notes,
                extra_source_cells=note_cells + extra_source_cells,
            )
        )
    return tuple(axes)


def _team_flags(axes: tuple[AxisExtract, ...]) -> tuple[str, ...]:
    if not axes:
        return ("NO_OPERATION_AXIS",)
    with_sources = sum(bool(axis.source_refs) for axis in axes)
    if with_sources == 0:
        return ("NO_SOURCE_REFERENCE",)
    if with_sources != len(axes):
        return ("SOME_AXES_NO_SOURCE_REFERENCE",)
    return ()


def _team_from_block(
    worksheet,
    media: MediaCatalog,
    *,
    team_id: str,
    display_order: int,
    row: int,
    first_column: int,
    operation_column: int,
    notes_first_column: int,
    notes_last_column: int,
    extra_source_column: int | None = None,
) -> TeamExtract | None:
    formation = _formation_block(
        worksheet,
        media,
        row=row,
        first_column=first_column,
    )
    if formation is None:
        return None
    members, last_row = formation
    axes = _axes_for_formation(
        worksheet,
        team_id=team_id,
        first_row=row,
        last_row=last_row,
        operation_column=operation_column,
        notes_first_column=notes_first_column,
        notes_last_column=notes_last_column,
        extra_source_column=extra_source_column,
    )
    notes = {axis.notes_raw for axis in axes}
    refs = _dedupe_refs(tuple(ref for axis in axes for ref in axis.source_refs))
    return TeamExtract(
        team_source_id=team_id,
        display_order=display_order,
        source_range=(
            f"{get_column_letter(first_column)}{row}:"
            f"{get_column_letter(notes_last_column)}{last_row}"
        ),
        formation=members,
        axes=axes,
        notes_raw=notes.pop() if len(notes) == 1 else None,
        source_refs=refs,
        flags=_team_flags(axes),
    )


def _deep_headers(worksheet, *, expected_area: int):
    headers = []
    for row in range(1, worksheet.max_row + 1):
        for column in (1, 2):
            cell = worksheet.cell(row, column)
            if cell.data_type == "f" or not isinstance(cell.value, str):
                continue
            normalized = unicodedata.normalize("NFKC", cell.value)
            match = _DEEP_STAGE_PATTERN.search(normalized)
            if match is None or int(match.group(1)) != expected_area:
                continue
            bounds = _merged_bounds(worksheet, row, column)
            wide_header = bounds[0] == row == bounds[2] and bounds[3] >= 19
            linked_header = (
                column == 2
                and worksheet.cell(row, 1).value is None
                and len(worksheet._images) > 0
            )
            if not wide_header and not linked_header:
                continue
            headers.append(
                (
                    int(match.group(2)),
                    row,
                    cell,
                    _range_label(bounds),
                )
            )
    by_stage = defaultdict(list)
    for item in headers:
        by_stage[item[0]].append(item)
    duplicates = {stage: items for stage, items in by_stage.items() if len(items) != 1}
    if duplicates:
        raise WorkbookLayoutError(
            f"{worksheet.title} has ambiguous deep stage headers: "
            + ", ".join(str(stage) for stage in sorted(duplicates))
        )
    return tuple(sorted((items[0] for items in by_stage.values()), key=lambda x: x[1]))


def _parse_deep_sheet(
    worksheet,
    media: MediaCatalog,
    *,
    element: str,
    area: int,
    strict: bool,
):
    headers = _deep_headers(worksheet, expected_area=area)
    if strict and {stage for stage, _row, _cell, _range in headers} != set(range(1, 11)):
        raise WorkbookLayoutError(
            f"{worksheet.title} does not contain exactly stages {area}-1 through {area}-10"
        )
    sections = []
    for header_index, (stage, header_row, header_cell, source_range) in enumerate(headers):
        last_row = (
            headers[header_index + 1][1] - 1
            if header_index + 1 < len(headers)
            else _last_content_row(worksheet)
        )
        panels = _DEEP_PANELS if stage < 10 else _DEEP_PANELS[:1]
        drafts = []
        for panel_index, (panel, first_col, operation_col, notes_first, notes_last) in enumerate(panels):
            for row in range(header_row + 1, last_row + 1):
                team_id = (
                    f"deep:{element.lower()}:{area:02d}-{stage:02d}:"
                    f"{panel.lower()}:{row:03d}"
                )
                team = _team_from_block(
                    worksheet,
                    media,
                    team_id=team_id,
                    display_order=0,
                    row=row,
                    first_column=first_col,
                    operation_column=operation_col,
                    notes_first_column=notes_first,
                    notes_last_column=(19 if stage == 10 else notes_last),
                    extra_source_column=(
                        _DEEP_LEFT_EXTRA_SOURCE_COLUMN if panel == "L" else None
                    ),
                )
                if team is not None:
                    drafts.append((panel_index, row, team))
        ordered = sorted(drafts, key=lambda item: (item[0], item[1]))
        teams = tuple(
            replace(team, display_order=index)
            for index, (_panel, _row, team) in enumerate(ordered, start=1)
        )
        if strict and not teams:
            raise WorkbookLayoutError(
                f"{worksheet.title} stage {area}-{stage} has no teams"
            )
        sections.append(
            SectionExtract(
                section_id=f"deep:{element.lower()}:{area:02d}-{stage:02d}",
                mode="DEEP",
                element=element,
                stage_refs=(StageRef(kind="DEEP", area=area, stage=stage),),
                label_raw=_cell_text(header_cell.value) or f"{area}-{stage}",
                source_range=source_range,
                era=EraMarker(kind="UNKNOWN"),
                teams=teams,
            )
        )
    return SheetExtract(
        sheet_name=worksheet.title,
        sheet_state=worksheet.sheet_state,
        layout_id="deep-8-10-two-panel-and-boss/v1",
        collection="DEEP_8_10_WORKBOOK",
        element=element,
        source_refs=_sheet_source_refs(worksheet),
        sections=tuple(sections),
    )


def _floor_group(value: Any) -> tuple[int, ...] | None:
    if not isinstance(value, str):
        return None
    normalized = unicodedata.normalize("NFKC", value)
    match = _FLOOR_PATTERN.search(normalized)
    if match is None:
        return None
    first = int(match.group(1))
    last = int(match.group(2) or first)
    if first < 1 or last > 10 or last < first:
        return None
    return tuple(range(first, last + 1))


def _remembrance_headers(worksheet):
    result = []
    for merged in worksheet.merged_cells.ranges:
        if not (
            merged.min_row == merged.max_row
            and merged.min_col == 2
            and merged.max_col == 14
        ):
            continue
        cell = worksheet.cell(merged.min_row, merged.min_col)
        floors = _floor_group(cell.value)
        if floors is not None:
            result.append((floors, merged.min_row, cell, str(merged)))
    by_floors = defaultdict(list)
    for item in result:
        by_floors[item[0]].append(item)
    if any(len(items) != 1 for items in by_floors.values()):
        raise WorkbookLayoutError(
            f"{worksheet.title} has duplicate remembrance floor headers"
        )
    return tuple(sorted(result, key=lambda item: item[1]))


def _parse_remembrance_sheet(
    worksheet,
    media: MediaCatalog,
    *,
    boss_key: str,
    expected_groups: tuple[tuple[int, ...], ...],
    strict: bool,
):
    headers = _remembrance_headers(worksheet)
    actual_groups = tuple(floors for floors, _row, _cell, _range in headers)
    if strict and actual_groups != expected_groups:
        raise WorkbookLayoutError(
            f"{worksheet.title} floor groups {actual_groups!r} do not match "
            f"the audited layout {expected_groups!r}"
        )
    sections = []
    for index, (floors, header_row, header_cell, source_range) in enumerate(headers):
        last_row = (
            headers[index + 1][1] - 1
            if index + 1 < len(headers)
            else _last_content_row(worksheet)
        )
        teams = []
        for row in range(header_row + 1, last_row + 1):
            floor_id = "-".join(str(floor) for floor in floors)
            team = _team_from_block(
                worksheet,
                media,
                team_id=f"remembrance:{boss_key.lower()}:{floor_id}:{row:03d}",
                display_order=len(teams) + 1,
                row=row,
                first_column=2,
                operation_column=7,
                notes_first_column=8,
                notes_last_column=14,
            )
            if team is not None:
                teams.append(team)
        if strict and not teams:
            raise WorkbookLayoutError(
                f"{worksheet.title} floor group {floors!r} has no teams"
            )
        sections.append(
            SectionExtract(
                section_id=(
                    f"remembrance:{boss_key.lower()}:"
                    + "-".join(f"{floor:02d}" for floor in floors)
                ),
                mode="REMEMBRANCE",
                element=boss_key,
                stage_refs=tuple(
                    StageRef(kind="REMEMBRANCE", area=1, stage=floor)
                    for floor in floors
                ),
                label_raw=_cell_text(header_cell.value) or str(floors),
                source_range=source_range,
                era=EraMarker(kind="UNKNOWN"),
                teams=tuple(teams),
            )
        )
    return SheetExtract(
        sheet_name=worksheet.title,
        sheet_state=worksheet.sheet_state,
        layout_id="remembrance-floor-groups/v1",
        collection="REMEMBRANCE_BATTLE",
        element=boss_key,
        source_refs=_sheet_source_refs(worksheet),
        sections=tuple(sections),
    )


def _luna_detail_team(
    worksheet,
    media: MediaCatalog,
    *,
    label: str,
    row: int,
    operation_coordinates: tuple[str, ...],
    note_coordinates: tuple[str, ...],
    display_order: int,
) -> TeamExtract | None:
    formation = _formation_block(
        worksheet,
        media,
        row=row,
        first_column=1,
    )
    if formation is None:
        return None
    members, _last_row = formation
    note_parts = []
    note_cells = []
    for coordinate in note_coordinates:
        cell = worksheet[coordinate]
        note_cells.append(cell)
        if cell.data_type != "f":
            raw = _cell_text(cell.value)
            if raw is not None and raw.strip():
                note_parts.append((cell, raw))
    axes = []
    for coordinate in operation_coordinates:
        operation_cell = worksheet[coordinate]
        axes.append(
            _make_axis(
                worksheet,
                axis_id=(
                    f"luna:top-ex:{label}:{row:03d}:axis:"
                    f"{operation_cell.coordinate.lower()}"
                ),
                operation_cell=operation_cell,
                note_parts=tuple(note_parts),
                extra_source_cells=tuple(note_cells),
            )
        )
    refs = _dedupe_refs(tuple(ref for axis in axes for ref in axis.source_refs))
    notes = "\n".join(raw for _cell, raw in note_parts) or None
    last_axis_row = max(worksheet[coordinate].row for coordinate in operation_coordinates)
    return TeamExtract(
        team_source_id=f"luna:top-ex:{label}:{row:03d}",
        display_order=display_order,
        source_range=f"A{row}:K{last_axis_row}",
        formation=members,
        axes=tuple(axes),
        notes_raw=notes,
        source_refs=refs,
        flags=_team_flags(tuple(axes)),
    )


def _parse_luna_sheet(worksheet, media: MediaCatalog, *, strict: bool):
    detail_teams = []
    for label, row, operations, notes in _LUNA_DETAIL_BLOCKS:
        team = _luna_detail_team(
            worksheet,
            media,
            label=label,
            row=row,
            operation_coordinates=operations,
            note_coordinates=notes,
            display_order=len(detail_teams) + 1,
        )
        if team is not None:
            detail_teams.append(team)
    if strict and len(detail_teams) != len(_LUNA_DETAIL_BLOCKS):
        raise WorkbookLayoutError(
            f"{worksheet.title} detailed strategy block has {len(detail_teams)} teams; "
            f"expected {len(_LUNA_DETAIL_BLOCKS)}"
        )

    second_header_row = None
    for merged in worksheet.merged_cells.ranges:
        if not (
            merged.min_row == merged.max_row
            and merged.min_col == 2
            and merged.max_col == 20
        ):
            continue
        text = _cell_text(worksheet.cell(merged.min_row, 2).value) or ""
        if "第2&3隊" in unicodedata.normalize("NFKC", text):
            second_header_row = merged.min_row
            break

    first_rows = []
    second_rows = []
    for row in range(1, _last_content_row(worksheet) + 1):
        # The prose-oriented detail area uses A:E.  Seen from the table's B:F
        # coordinates those rows contain four images, so only invoke the strict
        # formation parser after confirming all five table anchors are present.
        table_portraits = [
            media.image_sha(worksheet.title, row, column)
            for column in range(2, 7)
        ]
        if sum(digest is not None for digest in table_portraits) != 5:
            continue
        operation = _merged_anchor_cell(worksheet, row, 7)
        if second_header_row is not None and row > second_header_row:
            second_rows.append(row)
        elif second_header_row is None or row < second_header_row:
            if first_rows or (
                isinstance(operation.value, str)
                and "第1隊" in unicodedata.normalize("NFKC", operation.value)
            ):
                first_rows.append(row)

    def table_teams(rows: Iterable[int], label: str):
        teams = []
        for row in rows:
            team = _team_from_block(
                worksheet,
                media,
                team_id=f"luna:top-ex:{label}:{row:03d}",
                display_order=len(teams) + 1,
                row=row,
                first_column=2,
                operation_column=7,
                notes_first_column=8,
                notes_last_column=20,
            )
            if team is not None:
                teams.append(team)
        return tuple(teams)

    first_teams = table_teams(first_rows, "party-1-alternatives")
    second_teams = table_teams(second_rows, "party-2-3-alternatives")
    if strict and (len(first_teams), len(second_teams)) != (23, 20):
        raise WorkbookLayoutError(
            f"{worksheet.title} alternative tables have "
            f"{(len(first_teams), len(second_teams))} teams; expected (23, 20)"
        )

    stage_ref = (StageRef(kind="LUNA_TOWER_TOP_EX", area=1, stage=1),)
    sections = []
    if detail_teams or strict:
        sections.append(
            SectionExtract(
                section_id="luna:top-ex:detailed",
                mode="LUNA_TOWER",
                element="LUNA",
                stage_refs=stage_ref,
                label_raw=_cell_text(worksheet["A2"].value) or "露娜塔頂層EX",
                source_range="A2:T43",
                era=EraMarker(kind="UNKNOWN"),
                teams=tuple(detail_teams),
            )
        )
    if first_teams:
        sections.append(
            SectionExtract(
                section_id="luna:top-ex:party-1-alternatives",
                mode="LUNA_TOWER",
                element="LUNA",
                stage_refs=stage_ref,
                label_raw="第1隊",
                source_range=(
                    f"B{first_rows[0]}:T{first_rows[-1]}"
                    if first_rows
                    else "B1:T1"
                ),
                era=EraMarker(kind="UNKNOWN"),
                teams=first_teams,
            )
        )
    if second_teams:
        sections.append(
            SectionExtract(
                section_id="luna:top-ex:party-2-3-alternatives",
                mode="LUNA_TOWER",
                element="LUNA",
                stage_refs=stage_ref,
                label_raw=_cell_text(worksheet.cell(second_header_row, 2).value)
                if second_header_row is not None
                else "第2&3隊",
                source_range=(
                    f"B{second_header_row}:T{second_rows[-1]}"
                    if second_header_row is not None and second_rows
                    else "B1:T1"
                ),
                era=EraMarker(kind="UNKNOWN"),
                teams=second_teams,
            )
        )
    return SheetExtract(
        sheet_name=worksheet.title,
        sheet_state=worksheet.sheet_state,
        layout_id="luna-top-ex-detail-and-alternatives/v1",
        collection="LUNA_TOWER_TOP_EX",
        element="LUNA",
        source_refs=_sheet_source_refs(worksheet),
        sections=tuple(sections),
    )


def _consumed_occurrences(sheets: Iterable[SheetExtract]):
    consumed = set()
    for sheet in sheets:
        for section in sheet.sections:
            for team in section.teams:
                for member in team.formation:
                    column_letters = re.match(r"[A-Z]+", member.anchor_cell)
                    row_digits = re.search(r"\d+$", member.anchor_cell)
                    if column_letters is None or row_digits is None:
                        raise WorkbookLayoutError(
                            f"invalid member anchor {member.anchor_cell!r}"
                        )
                    column = 0
                    for character in column_letters.group(0):
                        column = column * 26 + ord(character) - ord("A") + 1
                    consumed.add(
                        (
                            sheet.sheet_name,
                            int(row_digits.group(0)),
                            column,
                            member.image_sha256,
                        )
                    )
    return consumed


def _unconsumed_portrait_runs(unconsumed, assets):
    portrait_digests = {
        asset.sha256
        for asset in assets
        if asset.width_px == 128 and asset.height_px == 128
    }
    columns_by_row = defaultdict(set)
    for sheet_name, row, column, digest in unconsumed:
        if digest in portrait_digests:
            columns_by_row[(sheet_name, row)].add(column)

    runs = set()
    for (sheet_name, row), columns in columns_by_row.items():
        ordered = sorted(columns)
        if not ordered:
            continue
        start = previous = ordered[0]
        for column in ordered[1:] + [None]:
            if column is not None and column == previous + 1:
                previous = column
                continue
            if previous - start + 1 >= 5:
                cells = tuple(
                    f"{get_column_letter(value)}{row}"
                    for value in range(start, previous + 1)
                )
                runs.add((sheet_name, cells))
            if column is not None:
                start = previous = column
    return frozenset(runs)


def _formula_count(worksheet) -> int:
    return sum(
        cell.data_type == "f"
        for row in worksheet.iter_rows()
        for cell in row
    )


def extract_workbook_2(
    path: str | Path,
    *,
    strict: bool = True,
) -> WorkbookExtract:
    input_path = Path(path)
    raw_bytes = input_path.read_bytes()
    workbook = load_workbook(BytesIO(raw_bytes), read_only=False, data_only=False)
    try:
        missing = [name for name in TARGET_SHEETS if name not in workbook.sheetnames]
        if missing:
            raise WorkbookLayoutError(
                "workbook is not the deep 8-10/remembrance/Luna source; missing sheets: "
                + ", ".join(missing)
            )
        if strict:
            hidden_targets = [
                name for name in TARGET_SHEETS if workbook[name].sheet_state != "visible"
            ]
            if hidden_targets:
                raise WorkbookLayoutError(
                    "target sheets unexpectedly hidden: " + ", ".join(hidden_targets)
                )

        workbook_states = {
            worksheet.title: worksheet.sheet_state for worksheet in workbook.worksheets
        }
        media = MediaCatalog.from_workbook(workbook, TARGET_SHEETS)
        sheets = [
            _parse_deep_sheet(
                workbook[name],
                media,
                element=element,
                area=area,
                strict=strict,
            )
            for name, element, area in DEEP_8_10_SHEETS
        ]
        sheets.extend(
            _parse_remembrance_sheet(
                workbook[name],
                media,
                boss_key=boss_key,
                expected_groups=expected_groups,
                strict=strict,
            )
            for name, boss_key, expected_groups in REMEMBRANCE_SHEETS
        )
        sheets.append(_parse_luna_sheet(workbook[LUNA_SHEET], media, strict=strict))
        assets = media.assets()

        diagnostics = []
        for sheet in sheets:
            for section in sheet.sections:
                for team in section.teams:
                    if not team.axes:
                        diagnostics.append(
                            {
                                "code": "TEAM_WITHOUT_OPERATION_AXIS",
                                "sheet_name": sheet.sheet_name,
                                "source_range": team.source_range,
                                "team_source_id": team.team_source_id,
                            }
                        )
        for name in TARGET_SHEETS:
            count = _formula_count(workbook[name])
            if count:
                diagnostics.append(
                    {
                        "code": "FORMULA_VALUES_NOT_USED_AS_DATA",
                        "formula_count": count,
                        "sheet_name": name,
                    }
                )
        target_occurrences = media.occurrence_keys()
        consumed = _consumed_occurrences(sheets)
        unexpected_consumed = consumed - target_occurrences
        if unexpected_consumed:
            raise WorkbookLayoutError(
                "parsed members do not match the embedded-media catalog"
            )
        unconsumed = target_occurrences - consumed
        unexpected_portrait_runs = (
            _unconsumed_portrait_runs(unconsumed, assets)
            - _AUDITED_UNCONSUMED_PORTRAIT_RUNS
        )
        if unexpected_portrait_runs:
            details = ", ".join(
                f"{sheet_name}!{cells[0]}:{cells[-1]}"
                for sheet_name, cells in sorted(unexpected_portrait_runs)
            )
            raise WorkbookLayoutError(
                "unconsumed portrait-like image runs may contain unsupported teams: "
                + details
            )
        by_sheet = defaultdict(list)
        for sheet_name, row, column, _digest in sorted(unconsumed):
            by_sheet[sheet_name].append(f"{get_column_letter(column)}{row}")
        for sheet_name in sorted(by_sheet):
            diagnostics.append(
                {
                    "anchor_cells": by_sheet[sheet_name],
                    "code": "UNCONSUMED_TARGET_SHEET_IMAGES",
                    "image_count": len(by_sheet[sheet_name]),
                    "sheet_name": sheet_name,
                }
            )
        for name in sorted(set(workbook.sheetnames) - set(TARGET_SHEETS)):
            state = workbook_states[name]
            diagnostics.append(
                {
                    "code": (
                        "HIDDEN_REFERENCE_SHEET"
                        if state in {"hidden", "veryHidden"}
                        else "SHEET_OUT_OF_SCOPE"
                    ),
                    "sheet_name": name,
                    "sheet_state": state,
                }
            )
    finally:
        workbook.close()

    section_count = sum(len(sheet.sections) for sheet in sheets)
    team_count = sum(
        len(section.teams) for sheet in sheets for section in sheet.sections
    )
    axis_count = sum(
        len(team.axes)
        for sheet in sheets
        for section in sheet.sections
        for team in section.teams
    )
    source_ref_count = sum(len(sheet.source_refs) for sheet in sheets) + sum(
        len(team.source_refs)
        for sheet in sheets
        for section in sheet.sections
        for team in section.teams
    )
    summary = {
        "asset_count": len(assets),
        "asset_occurrence_count": sum(len(asset.occurrences) for asset in assets),
        "axis_count": axis_count,
        "axis_counts_by_sheet": {
            sheet.sheet_name: sum(
                len(team.axes) for section in sheet.sections for team in section.teams
            )
            for sheet in sheets
        },
        "formula_ignored_count": sum(
            diagnostic.get("formula_count", 0) for diagnostic in diagnostics
        ),
        "section_count": section_count,
        "sheet_count": len(sheets),
        "source_ref_count": source_ref_count,
        "team_count": team_count,
        "team_counts_by_sheet": {
            sheet.sheet_name: sum(len(section.teams) for section in sheet.sections)
            for sheet in sheets
        },
        "unconsumed_image_count": sum(
            diagnostic.get("image_count", 0) for diagnostic in diagnostics
        ),
        "unmapped_member_count": team_count * 5,
    }
    return WorkbookExtract(
        schema_version=SCHEMA_VERSION,
        parser={
            "layout_ids": [
                "deep-8-10-two-panel-and-boss/v1",
                "remembrance-floor-groups/v1",
                "luna-top-ex-detail-and-alternatives/v1",
            ],
            "name": PARSER_NAME,
            "version": PARSER_VERSION,
        },
        source_workbook={
            "byte_length": len(raw_bytes),
            "filename": input_path.name,
            "sha256": hashlib.sha256(raw_bytes).hexdigest(),
        },
        sheets=tuple(sheets),
        assets=assets,
        summary=summary,
        diagnostics=tuple(
            sorted(
                diagnostics,
                key=lambda item: (item["code"], item["sheet_name"]),
            )
        ),
    )
