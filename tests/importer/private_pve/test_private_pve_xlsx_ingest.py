from __future__ import annotations

from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import PatternFill
from PIL import Image

from pcr_pipeline.private_pve.cli import format_stdout_summary, main
from pcr_pipeline.private_pve.deep_1_7 import extract_deep_1_7_workbook
from pcr_pipeline.private_pve.media import MediaCatalog
from pcr_pipeline.private_pve.models import WorkbookLayoutError
from pcr_pipeline.private_pve.operations import (
    merge_parsed_operations,
    parse_operation,
)
from pcr_pipeline.private_pve.sources import SOURCE_ALIASES, extract_source_refs


DEEP_SHEETS = (
    "紅焔の深域(火屬性)",
    "蒼波の深域(水屬性)",
    "翠嵐の深域(風屬性)",
    "珀天の深域(光屬性)",
    "紫冥の深域(闇屬性)",
)
ANEMONE_SHEET = "安涅默涅道中通關組合"


def _portrait_png() -> bytes:
    payload = BytesIO()
    Image.new("RGBA", (8, 8), (231, 76, 60, 255)).save(payload, format="PNG")
    return payload.getvalue()


def _add_formation(worksheet, *, row: int, first_column: int, payload: bytes) -> None:
    for column in range(first_column, first_column + 5):
        portrait = WorksheetImage(BytesIO(payload))
        portrait.anchor = worksheet.cell(row=row, column=column).coordinate
        worksheet.add_image(portrait)


def _build_partial_workbook(path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    portrait = _portrait_png()

    for sheet_name in DEEP_SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.merge_cells("A1:I1")
        worksheet["A1"] = "　１－１　"
        _add_formation(worksheet, row=2, first_column=1, payload=portrait)
        worksheet["F2"] = "XOOOO"
        worksheet["I2"] = "https://example.invalid/first"

    fire = workbook[DEEP_SHEETS[0]]
    fire["F2"] = "半自動"
    fire["G2"] = "XOOOO"
    _add_formation(fire, row=3, first_column=1, payload=portrait)
    fire["F3"] = "全SET / OOOXO"
    fire["I3"] = "#2"

    workbook[DEEP_SHEETS[-1]].sheet_state = "hidden"

    anemone = workbook.create_sheet(ANEMONE_SHEET)
    anemone["A3"] = datetime(2024, 1, 1)
    anemone["A3"].number_format = "m-d"
    anemone["A3"].fill = PatternFill(
        fill_type="solid",
        fgColor="FFEA9999",
    )
    _add_formation(anemone, row=3, first_column=2, payload=portrait)
    anemone["G3"] = "全SET"
    anemone["I3"] = "#4"

    workbook.save(path)
    workbook.close()


def test_operation_patterns_preserve_slot_order_and_variants() -> None:
    operation = parse_operation("全SET / XOOOO")

    assert operation.order_basis == "WORKBOOK_TEXT_LEFT_TO_RIGHT"
    assert operation.member_alignment == "UNRESOLVED"
    assert [variant.kind for variant in operation.variants] == [
        "FULL_SET",
        "POSITIONAL_PATTERN",
    ]
    assert operation.variants[0].source_order_states == ("SET",) * 5
    assert operation.variants[1].source_order_states == (
        "NOT_SET",
        "SET",
        "SET",
        "SET",
        "SET",
    )


def test_operation_tokens_in_notes_keep_their_field_origin() -> None:
    operation = merge_parsed_operations(
        parse_operation(
            "半自動",
            origin_field="OPERATION_TEXT",
            origin_cell="F23",
        ),
        parse_operation(
            "XOOOO",
            origin_field="NOTES_TEXT",
            origin_cell="G23",
        ),
    )

    assert operation.kind == "COMPOSITE"
    assert operation.execution_hints == ("SEMI_AUTO",)
    assert operation.variants[0].origin_field == "NOTES_TEXT"
    assert operation.variants[0].origin_cell == "G23"
    assert operation.variants[0].source_order_states[0] == "NOT_SET"


@pytest.mark.parametrize("alias_id", sorted(SOURCE_ALIASES))
def test_source_aliases_are_cell_local(alias_id: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = "https://example.invalid/previous"
    worksheet["A2"] = alias_id

    first = extract_source_refs(worksheet["A1"], sheet_name="fixture")
    second = extract_source_refs(worksheet["A2"], sheet_name="fixture")

    assert [source.kind for source in first] == ["URL"]
    assert [source.alias_id for source in second] == [alias_id]
    assert all(source.url is None for source in second)
    workbook.close()


def test_hyperlink_formula_extracts_only_the_first_argument() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet["A1"] = '=HYPERLINK("https://youtu.be/abc","影片")'

    refs = extract_source_refs(worksheet["A1"], sheet_name="fixture")

    assert len(refs) == 1
    assert refs[0].url == "https://youtu.be/abc"
    workbook.close()


def test_partial_fixture_is_deterministic_and_deduplicates_images(tmp_path) -> None:
    source = tmp_path / "fixture.xlsx"
    _build_partial_workbook(source)

    first = extract_deep_1_7_workbook(source, strict=False)
    second = extract_deep_1_7_workbook(source, strict=False)

    assert first.canonical_json_bytes() == second.canonical_json_bytes()
    assert first.summary["team_count"] == 6
    assert first.summary["axis_count"] == 7
    assert first.summary["asset_count"] == 1
    assert first.summary["asset_occurrence_count"] == 35

    fire = first.sheets[0]
    assert fire.sections[0].stage_refs[0].area == 1
    assert fire.sections[0].stage_refs[0].stage == 1
    assert len(fire.sections[0].teams) == 1
    assert len(fire.sections[0].teams[0].axes) == 2
    note_variant = fire.sections[0].teams[0].axes[0].operation.variants[0]
    assert note_variant.origin_field == "NOTES_TEXT"
    assert note_variant.source_order_states[0] == "NOT_SET"
    alias_refs = [
        source
        for source in fire.sections[0].teams[0].source_refs
        if source.alias_id == "#2"
    ]
    assert len(alias_refs) == 1
    assert alias_refs[0].url is None

    anemone = first.sheets[-1]
    assert anemone.sections[0].stage_refs[0].area == 1
    assert anemone.sections[0].stage_refs[0].stage == 1
    assert anemone.sections[0].label_raw == "1-1"
    assert first.sheets[-2].sheet_state == "hidden"


def test_merged_source_cell_is_inherited_with_explicit_provenance(tmp_path) -> None:
    source = tmp_path / "fixture.xlsx"
    _build_partial_workbook(source)
    workbook = load_workbook(source)
    fire = workbook[DEEP_SHEETS[0]]
    fire["I3"] = None
    fire.merge_cells("I2:I3")
    workbook.save(source)
    workbook.close()

    result = extract_deep_1_7_workbook(source, strict=False)
    team = result.sheets[0].sections[0].teams[0]

    assert len(team.axes) == 2
    inherited = team.axes[1].source_refs[0]
    assert inherited.origin_cell.endswith("!I2")
    assert inherited.applies_to_cell == "I3"


def test_media_catalog_refuses_implicit_image_conversion() -> None:
    class UnsupportedImage:
        format = "webp"
        anchor = "A1"

        def _data(self):  # pragma: no cover - must fail before conversion
            raise AssertionError("unsupported image must not be converted")

    class Worksheet:
        _images = [UnsupportedImage()]

    class FakeWorkbook:
        def __getitem__(self, _sheet_name):
            return Worksheet()

    with pytest.raises(WorkbookLayoutError, match="implicit Pillow conversion"):
        MediaCatalog.from_workbook(FakeWorkbook(), ["fixture"])


def test_cli_writes_atomically_and_prints_ascii_safe_summary(tmp_path, capsys) -> None:
    source = tmp_path / "fixture.xlsx"
    destination = tmp_path / "nested" / "staging.json"
    _build_partial_workbook(source)

    assert main(
        [
            "extract",
            "--input",
            str(source),
            "--output",
            str(destination),
            "--allow-partial-layout",
        ]
    ) == 0

    expected = extract_deep_1_7_workbook(source, strict=False)
    assert destination.read_bytes() == expected.canonical_json_bytes()
    assert not list(destination.parent.glob(f".{destination.name}.*.tmp"))
    status_line = capsys.readouterr().out.strip()
    assert status_line == format_stdout_summary(expected.summary)
    assert status_line.isascii()
    status_line.encode("cp950")


def test_wrong_workbook_fails_closed(tmp_path) -> None:
    source = tmp_path / "wrong.xlsx"
    workbook = Workbook()
    workbook.save(source)
    workbook.close()

    with pytest.raises(WorkbookLayoutError, match="missing sheets"):
        extract_deep_1_7_workbook(source)


def test_cli_refuses_to_overwrite_the_source_workbook(tmp_path) -> None:
    source = tmp_path / "fixture.xlsx"
    original = b"not a workbook"
    source.write_bytes(original)

    with pytest.raises(SystemExit):
        main(["extract", "--input", str(source), "--output", str(source)])

    assert source.read_bytes() == original


def test_cli_requires_json_output(tmp_path) -> None:
    source = tmp_path / "fixture.xlsx"
    _build_partial_workbook(source)
    destination = tmp_path / "another-workbook.xlsx"

    with pytest.raises(SystemExit):
        main(
            [
                "extract",
                "--input",
                str(source),
                "--output",
                str(destination),
                "--allow-partial-layout",
            ]
        )

    assert not destination.exists()
