from __future__ import annotations

import copy
import hashlib
import json
from dataclasses import asdict
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image

from pcr_pipeline.xlsx_ingest.catalog import (
    LOCAL_CATALOG_SCHEMA_VERSION,
    CatalogMergeError,
    export_staging_media,
    merge_staging_catalogs,
)
from pcr_pipeline.xlsx_ingest.cli import main
from pcr_pipeline.xlsx_ingest.media import MediaCatalog
from pcr_pipeline.xlsx_ingest.models import SCHEMA_VERSION


def _portrait_png() -> bytes:
    payload = BytesIO()
    Image.new("RGBA", (8, 8), (155, 89, 182, 255)).save(payload, format="PNG")
    return payload.getvalue()


def _build_workbook(path, *, sheet_name: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    portrait = WorksheetImage(BytesIO(_portrait_png()))
    portrait.anchor = "A1"
    worksheet.add_image(portrait)
    workbook.save(path)
    workbook.close()


def _staging_document(path, *, sheet_name: str, diagnostic_code: str):
    raw_bytes = path.read_bytes()
    workbook = load_workbook(BytesIO(raw_bytes), data_only=False, read_only=False)
    try:
        media = MediaCatalog.from_workbook(workbook, [sheet_name])
        assets = json.loads(
            json.dumps([asdict(asset) for asset in media.assets()])
        )
    finally:
        workbook.close()
    digest = assets[0]["sha256"]
    return {
        "assets": assets,
        "diagnostics": [{"code": diagnostic_code, "sheet_name": sheet_name}],
        "parser": {
            "layout_ids": ["fixture/v1"],
            "name": "fixture-parser",
            "version": "1.0.0",
        },
        "schema_version": SCHEMA_VERSION,
        "sheets": [
            {
                "sections": [
                    {
                        "teams": [
                            {
                                "axes": [],
                                "formation": [{"image_sha256": digest}],
                            }
                        ]
                    }
                ],
                "sheet_name": sheet_name,
            }
        ],
        "source_workbook": {
            "byte_length": len(raw_bytes),
            "filename": path.name,
            "sha256": hashlib.sha256(raw_bytes).hexdigest(),
        },
        "summary": {"asset_count": 1, "sheet_count": 1},
    }


def _write_staging(path, document) -> None:
    path.write_text(
        json.dumps(document, ensure_ascii=False, separators=(",", ":"), sort_keys=True)
        + "\n",
        encoding="utf-8",
    )


def test_media_export_is_exact_atomic_idempotent_and_fails_on_conflict(
    tmp_path,
) -> None:
    workbook_path = tmp_path / "source.xlsx"
    catalog_path = tmp_path / "source.json"
    output_directory = tmp_path / ".runtime" / "private_pve" / "assets"
    _build_workbook(workbook_path, sheet_name="來源")
    staging = _staging_document(
        workbook_path,
        sheet_name="來源",
        diagnostic_code="FIRST_SOURCE",
    )
    _write_staging(catalog_path, staging)

    first = export_staging_media(
        workbook_path=workbook_path,
        staging_catalog_path=catalog_path,
        output_directory=output_directory,
    )
    asset = staging["assets"][0]
    exported = output_directory / f"{asset['sha256']}.{asset['file_extension']}"

    assert first["written_count"] == 1
    assert first["verified_existing_count"] == 0
    assert hashlib.sha256(exported.read_bytes()).hexdigest() == asset["sha256"]
    assert not list(output_directory.glob(".*.tmp"))

    second = export_staging_media(
        workbook_path=workbook_path,
        staging_catalog_path=catalog_path,
        output_directory=output_directory,
    )
    assert second["written_count"] == 0
    assert second["verified_existing_count"] == 1

    exported.write_bytes(b"conflicting existing bytes")
    with pytest.raises(CatalogMergeError, match="existing exported asset"):
        export_staging_media(
            workbook_path=workbook_path,
            staging_catalog_path=catalog_path,
            output_directory=output_directory,
        )
    assert exported.read_bytes() == b"conflicting existing bytes"


def test_media_export_requires_runtime_and_matching_workbook_digest(tmp_path) -> None:
    workbook_path = tmp_path / "source.xlsx"
    catalog_path = tmp_path / "source.json"
    _build_workbook(workbook_path, sheet_name="來源")
    staging = _staging_document(
        workbook_path,
        sheet_name="來源",
        diagnostic_code="FIRST_SOURCE",
    )
    _write_staging(catalog_path, staging)

    with pytest.raises(CatalogMergeError, match="directory named .runtime"):
        export_staging_media(
            workbook_path=workbook_path,
            staging_catalog_path=catalog_path,
            output_directory=tmp_path / "assets",
        )

    changed = copy.deepcopy(staging)
    changed["source_workbook"]["sha256"] = "0" * 64
    _write_staging(catalog_path, changed)
    with pytest.raises(CatalogMergeError, match="provenance"):
        export_staging_media(
            workbook_path=workbook_path,
            staging_catalog_path=catalog_path,
            output_directory=tmp_path / ".runtime" / "assets",
        )


def test_merge_is_order_independent_and_preserves_source_scoped_data(tmp_path) -> None:
    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"
    _build_workbook(first_path, sheet_name="第一來源")
    _build_workbook(second_path, sheet_name="第二來源")
    first = _staging_document(
        first_path,
        sheet_name="第一來源",
        diagnostic_code="FIRST_SOURCE",
    )
    second = _staging_document(
        second_path,
        sheet_name="第二來源",
        diagnostic_code="SECOND_SOURCE",
    )

    forward = merge_staging_catalogs([first, second])
    reverse = merge_staging_catalogs([second, first])

    assert forward.schema_version == LOCAL_CATALOG_SCHEMA_VERSION
    assert forward.canonical_json_bytes() == reverse.canonical_json_bytes()
    assert forward.summary == {
        "asset_count": 1,
        "asset_occurrence_count": 2,
        "axis_count": 0,
        "diagnostic_count": 2,
        "section_count": 2,
        "sheet_count": 2,
        "source_workbook_count": 2,
        "team_count": 2,
    }
    assert {
        source["diagnostics"][0]["code"] for source in forward.sources
    } == {"FIRST_SOURCE", "SECOND_SOURCE"}
    assert {
        source["sheets"][0]["sheet_name"] for source in forward.sources
    } == {"第一來源", "第二來源"}
    assert len(forward.assets) == 1
    assert len(forward.assets[0]["occurrences"]) == 2
    assert all(
        occurrence["source_workbook_sha256"]
        for occurrence in forward.assets[0]["occurrences"]
    )


def test_merge_rejects_schema_digest_and_metadata_conflicts(tmp_path) -> None:
    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"
    _build_workbook(first_path, sheet_name="第一來源")
    _build_workbook(second_path, sheet_name="第二來源")
    first = _staging_document(
        first_path,
        sheet_name="第一來源",
        diagnostic_code="FIRST_SOURCE",
    )
    second = _staging_document(
        second_path,
        sheet_name="第二來源",
        diagnostic_code="SECOND_SOURCE",
    )

    wrong_schema = copy.deepcopy(second)
    wrong_schema["schema_version"] = "private-pve-xlsx-staging/v999"
    with pytest.raises(CatalogMergeError, match="schema must be"):
        merge_staging_catalogs([first, wrong_schema])

    wrong_digest = copy.deepcopy(second)
    wrong_digest["assets"][0]["sha256"] = "not-a-digest"
    with pytest.raises(CatalogMergeError, match="lowercase SHA-256"):
        merge_staging_catalogs([first, wrong_digest])

    conflicting_metadata = copy.deepcopy(second)
    conflicting_metadata["assets"][0]["width_px"] += 1
    with pytest.raises(CatalogMergeError, match="conflicting metadata"):
        merge_staging_catalogs([first, conflicting_metadata])

    conflicting_source = copy.deepcopy(first)
    conflicting_source["diagnostics"].append({"code": "TAMPERED"})
    with pytest.raises(CatalogMergeError, match="claim source workbook digest"):
        merge_staging_catalogs([first, conflicting_source])


def test_merge_cli_enforces_json_paths_and_does_not_overwrite_inputs(
    tmp_path,
    capsys,
) -> None:
    first_path = tmp_path / "first.xlsx"
    second_path = tmp_path / "second.xlsx"
    first_json = tmp_path / "first.json"
    second_json = tmp_path / "second.json"
    output = tmp_path / ".runtime" / "private_pve" / "catalog.json"
    _build_workbook(first_path, sheet_name="第一來源")
    _build_workbook(second_path, sheet_name="第二來源")
    _write_staging(
        first_json,
        _staging_document(
            first_path,
            sheet_name="第一來源",
            diagnostic_code="FIRST_SOURCE",
        ),
    )
    _write_staging(
        second_json,
        _staging_document(
            second_path,
            sheet_name="第二來源",
            diagnostic_code="SECOND_SOURCE",
        ),
    )

    assert main(
        [
            "merge-catalogs",
            "--input",
            str(second_json),
            "--input",
            str(first_json),
            "--output",
            str(output),
        ]
    ) == 0
    assert json.loads(output.read_text(encoding="utf-8"))["schema_version"] == (
        LOCAL_CATALOG_SCHEMA_VERSION
    )
    assert capsys.readouterr().out.strip().isascii()

    original = first_json.read_bytes()
    with pytest.raises(SystemExit):
        main(
            [
                "merge-catalogs",
                "--input",
                str(first_json),
                "--input",
                str(second_json),
                "--output",
                str(first_json),
            ]
        )
    assert first_json.read_bytes() == original

    with pytest.raises(SystemExit):
        main(
            [
                "merge-catalogs",
                "--input",
                str(first_json),
                "--input",
                str(second_json.with_suffix(".txt")),
                "--output",
                str(output),
            ]
        )
