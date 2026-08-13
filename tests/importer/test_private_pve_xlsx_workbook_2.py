from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as WorksheetImage
from PIL import Image

from pcr_pipeline.xlsx_ingest.cli import main
from pcr_pipeline.xlsx_ingest.models import WorkbookLayoutError
from pcr_pipeline.xlsx_ingest.workbook_2 import (
    DEEP_8_10_SHEETS,
    LUNA_SHEET,
    REMEMBRANCE_SHEETS,
    extract_workbook_2,
)


def _portrait_png() -> bytes:
    payload = BytesIO()
    Image.new("RGBA", (128, 128), (52, 152, 219, 255)).save(payload, format="PNG")
    return payload.getvalue()


def _add_formation(worksheet, *, row: int, first_column: int, payload: bytes) -> None:
    for column in range(first_column, first_column + 5):
        portrait = WorksheetImage(BytesIO(payload))
        portrait.anchor = worksheet.cell(row=row, column=column).coordinate
        worksheet.add_image(portrait)


def _build_partial_workbook_2(path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    portrait = _portrait_png()

    for sheet_name, _element, area in DEEP_8_10_SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        worksheet.merge_cells("A1:S1")
        worksheet["A1"] = f" {area} - 1"
        _add_formation(worksheet, row=2, first_column=2, payload=portrait)
        worksheet["G2"] = "全SET"
        worksheet["H2"] = "TP test"

    fire = workbook["火8"]
    for column in range(2, 7):
        fire.merge_cells(
            start_row=2,
            start_column=column,
            end_row=3,
            end_column=column,
        )
    fire["G2"] = "半A"
    fire["G2"].hyperlink = "https://example.invalid/axis-1"
    fire["H2"] = "XOOOO"
    fire["I2"] = '=HYPERLINK("https://example.invalid/note","影片")'
    fire["J2"] = "全SET替代"
    fire["J2"].hyperlink = "https://example.invalid/full-set-video"
    fire["G3"] = "全SET"
    fire["H3"] = "OOOXO"

    # Formula numbering is intentionally present but must not become identity
    # or operation data.
    workbook["水8"]["A2"] = "=1+0"

    for sheet_name, _boss_key, expected_groups in REMEMBRANCE_SHEETS:
        worksheet = workbook.create_sheet(sheet_name)
        floors = expected_groups[0]
        label = (
            f" {floors[0]}-{floors[-1]}層"
            if len(floors) > 1
            else f" {floors[0]}層"
        )
        worksheet.merge_cells("B1:N1")
        worksheet["B1"] = label
        _add_formation(worksheet, row=2, first_column=2, payload=portrait)
        worksheet["G2"] = "OXOOO"
        worksheet["H2"] = "備註"

    luna = workbook.create_sheet(LUNA_SHEET)
    _add_formation(luna, row=1, first_column=2, payload=portrait)
    luna["G1"] = "第1隊"
    luna.merge_cells("H1:T1")
    luna["H1"] = "全SET / XOOOO"

    reference = workbook.create_sheet("EX裝備需求")
    reference.sheet_state = "hidden"
    reference["A1"] = "reference only"

    workbook.save(path)
    workbook.close()


def test_workbook_2_partial_fixture_is_deterministic_and_preserves_axes(tmp_path) -> None:
    source = tmp_path / "workbook-2.xlsx"
    _build_partial_workbook_2(source)

    first = extract_workbook_2(source, strict=False)
    second = extract_workbook_2(source, strict=False)

    assert first.canonical_json_bytes() == second.canonical_json_bytes()
    assert first.summary["sheet_count"] == 20
    assert first.summary["section_count"] == 20
    assert first.summary["team_count"] == 20
    assert first.summary["axis_count"] == 21
    assert first.summary["asset_count"] == 1
    assert first.summary["asset_occurrence_count"] == 100
    assert first.summary["unconsumed_image_count"] == 0

    fire = first.sheets[0].sections[0].teams[0]
    assert len(fire.axes) == 2
    assert fire.axes[0].operation.execution_hints == ("SEMI_AUTO",)
    assert fire.axes[0].operation.variants[0].origin_field == "NOTES_TEXT"
    assert fire.axes[0].operation.variants[0].source_order_states[0] == "NOT_SET"
    assert {ref.url for ref in fire.axes[0].source_refs} == {
        "https://example.invalid/axis-1",
        "https://example.invalid/full-set-video",
        "https://example.invalid/note",
    }
    assert any(
        diagnostic["code"] == "HIDDEN_REFERENCE_SHEET"
        and diagnostic["sheet_name"] == "EX裝備需求"
        for diagnostic in first.diagnostics
    )
    assert first.summary["formula_ignored_count"] == 2


def test_workbook_2_cli_uses_separate_atomic_command(tmp_path, capsys) -> None:
    source = tmp_path / "workbook-2.xlsx"
    destination = tmp_path / "staging.json"
    _build_partial_workbook_2(source)

    assert main(
        [
            "extract-workbook-2",
            "--input",
            str(source),
            "--output",
            str(destination),
            "--allow-partial-layout",
        ]
    ) == 0

    expected = extract_workbook_2(source, strict=False)
    assert destination.read_bytes() == expected.canonical_json_bytes()
    assert capsys.readouterr().out.strip().isascii()


def test_workbook_2_formula_in_operation_fails_closed(tmp_path) -> None:
    source = tmp_path / "workbook-2.xlsx"
    _build_partial_workbook_2(source)
    workbook = load_workbook(source)
    workbook["水8"]["G2"] = "=1+1"
    workbook.save(source)
    workbook.close()

    with pytest.raises(WorkbookLayoutError, match="operation is a formula"):
        extract_workbook_2(source, strict=False)


def test_workbook_2_partial_roster_fails_closed(tmp_path) -> None:
    source = tmp_path / "workbook-2.xlsx"
    _build_partial_workbook_2(source)
    workbook = load_workbook(source)
    worksheet = workbook["風8"]
    portrait = WorksheetImage(BytesIO(_portrait_png()))
    portrait.anchor = "B3"
    worksheet.add_image(portrait)
    workbook.save(source)
    workbook.close()

    with pytest.raises(WorkbookLayoutError, match="expected exactly 5"):
        extract_workbook_2(source, strict=False)


def test_workbook_2_wrong_source_fails_closed(tmp_path) -> None:
    source = tmp_path / "wrong.xlsx"
    workbook = Workbook()
    workbook.save(source)
    workbook.close()

    with pytest.raises(WorkbookLayoutError, match="missing sheets"):
        extract_workbook_2(source)


def test_workbook_2_unconsumed_portrait_run_fails_closed(tmp_path) -> None:
    source = tmp_path / "workbook-2.xlsx"
    _build_partial_workbook_2(source)
    workbook = load_workbook(source)
    worksheet = workbook["火8"]
    worksheet.merge_cells("A4:S4")
    worksheet["A4"] = " 8 - 10"
    _add_formation(
        worksheet,
        row=5,
        first_column=12,
        payload=_portrait_png(),
    )
    worksheet["Q5"] = "全SET"
    workbook.save(source)
    workbook.close()

    with pytest.raises(
        WorkbookLayoutError,
        match="unconsumed portrait-like image runs",
    ):
        extract_workbook_2(source, strict=False)
